"""Shared helpers for the case-deck scripts.

Everything that touches the PowerPoint package lives here, so build, edit,
navigate and review all clone, fill and drop slides the same way.
"""
from __future__ import annotations

import copy
import hashlib
import io
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.util import Emu, Pt

SKILL_DIR = Path(__file__).resolve().parent.parent
ASSETS = SKILL_DIR / "assets"
TEMPLATE = ASSETS / "template.pptx"
PATTERNS = ASSETS / "patterns.pptx"
CATALOG = ASSETS / "catalog.yaml"

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "pslz": "http://schemas.microsoft.com/office/powerpoint/2016/slidezoom",
    "p166": "http://schemas.microsoft.com/office/powerpoint/2016/6/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}
R_NS = NS["r"]

# House palette (RGB hex, upper case). Theme colours are always allowed.
PALETTE = {
    "061F32", "061E31", "22A3DF", "3478AB", "1D4769", "2222F6", "F8BC3B",
    "F2F2F2", "FFFFFF", "000000", "D9D9D9", "BFBFBF", "A6A6A6", "7F7F7F",
    "595959", "404040", "262626", "E7E6E6", "EDEDED", "F7F7F7",
}

LEFTOVER = re.compile(
    r"(\[Insert|\[Title|Lorem|Click to add|\bTBD\b|\bxx\b|Company Name|\[Section\]|\[Subtitle\]|"
    r"Findings and recommendations summary|Click to edit)",
    re.I,
)


# --------------------------------------------------------------------------
# Opening, saving, slide lists
# --------------------------------------------------------------------------

def open_deck(path) -> Presentation:
    return Presentation(str(path))


def sld_id_list(prs):
    return prs.slides._sldIdLst


def slide_index(prs, slide) -> int:
    """0-based position of a slide in the deck."""
    for i, s in enumerate(prs.slides):
        if s.slide_id == slide.slide_id:
            return i
    raise ValueError("slide not in deck")


def drop_slide(prs, slide):
    lst = sld_id_list(prs)
    for sid in list(lst):
        if int(sid.get("id")) == slide.slide_id:
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
            return
    raise ValueError("slide not found")


def keep_only(prs, keep_positions):
    """Drop every slide whose 0-based position is not in keep_positions."""
    keep = set(keep_positions)
    for i, s in reversed(list(enumerate(prs.slides))):
        if i not in keep:
            drop_slide(prs, s)


def move_slide(prs, slide, new_pos: int):
    lst = sld_id_list(prs)
    for sid in list(lst):
        if int(sid.get("id")) == slide.slide_id:
            lst.remove(sid)
            lst.insert(new_pos, sid)
            return
    raise ValueError("slide not found")


def layout_by_name(prs, name, fallbacks=("Default", "Blank")):
    for n in (name, *fallbacks):
        for master in prs.slide_masters:
            for lay in master.slide_layouts:
                if lay.name == n:
                    return lay
    return prs.slide_layouts[0]


# --------------------------------------------------------------------------
# Cloning slides (same package or across packages)
# --------------------------------------------------------------------------

_SKIP_RELS = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}


def _partname_template(partname: str) -> str:
    m = re.match(r"^(.*?)(\d+)(\.[A-Za-z0-9]+)$", partname)
    if m:
        return f"{m.group(1)}%d{m.group(3)}"
    stem, dot, ext = partname.rpartition(".")
    return f"{stem}%d.{ext}"


def _copy_part(part, dst_package, seen):
    """Deep-copy a non-image part (chart, embedding, ...) into dst_package."""
    if id(part) in seen:
        return seen[id(part)]
    new_name = dst_package.next_partname(_partname_template(str(part.partname)))
    new = part.__class__.load(PackURI(new_name), part.content_type, dst_package, part.blob)
    seen[id(part)] = new
    mapping = {}
    for rId, rel in list(part.rels.items()):
        if rel.is_external:
            mapping[rId] = new.relate_to(rel.target_ref, rel.reltype, is_external=True)
            continue
        child = _copy_part(rel.target_part, dst_package, seen)
        mapping[rId] = new.relate_to(child, rel.reltype)
    _rewrite_rids(getattr(new, "_element", None), mapping)
    return new


def _rewrite_rids(element, mapping):
    """Rewrite r:* attributes in one pass (no chained replacements)."""
    if element is None or not mapping:
        return
    for el in element.iter():
        for attr, val in list(el.attrib.items()):
            if attr.startswith("{%s}" % R_NS) and val in mapping:
                el.set(attr, mapping[val])


def _ensure_unique_partname(prs, part):
    """python-pptx names a new slide slide<count+1>.xml, which can clash
    with an existing part after slides were dropped. Rename if needed."""
    used = {str(p.partname) for p in prs.part.package.iter_parts() if p is not part}
    if str(part.partname) not in used:
        return
    i = 1
    while f"/ppt/slides/slide{i}.xml" in used:
        i += 1
    part.partname = PackURI(f"/ppt/slides/slide{i}.xml")


def clone_slide(src_slide, dst_prs, layout_name=None, drop_internal_links=True):
    """Append a copy of src_slide to dst_prs and return the new slide.

    Works inside one package and across packages. Pictures are re-embedded
    (deduplicated by hash), charts and other parts are copied, external
    hyperlinks are kept, links to other slides are removed.
    """
    want = layout_name or src_slide.slide_layout.name
    lay = layout_by_name(dst_prs, want)
    foreign = lay.name != want
    new = dst_prs.slides.add_slide(lay)
    _ensure_unique_partname(dst_prs, new.part)
    src_el, dst_el = src_slide._element, new._element

    mapping = {}
    seen = {}
    for rId, rel in src_slide.part.rels.items():
        if rel.reltype in _SKIP_RELS:
            continue
        if rel.is_external:
            mapping[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            continue
        if rel.reltype == RT.SLIDE:
            mapping[rId] = None  # internal jump; dropped below
            continue
        target = rel.target_part
        if rel.reltype == RT.IMAGE:
            try:
                _, new_rId = new.part.get_or_add_image_part(io.BytesIO(target.blob))
            except Exception:  # SVG, EMF and other formats Pillow cannot read
                new_rId = new.part.relate_to(_copy_part(target, new.part.package, seen), rel.reltype)
            mapping[rId] = new_rId
        else:
            mapping[rId] = new.part.relate_to(_copy_part(target, new.part.package, seen), rel.reltype)

    # copy the slide body in place (python-pptx caches the shape tree object,
    # so the existing cSld/spTree elements must stay and only their children change)
    p = NS["p"]
    d_csld, s_csld = dst_el.find(f"{{{p}}}cSld"), src_el.find(f"{{{p}}}cSld")
    d_tree, s_tree = d_csld.find(f"{{{p}}}spTree"), s_csld.find(f"{{{p}}}spTree")
    for child in list(d_tree):
        d_tree.remove(child)
    for child in s_tree:
        d_tree.append(copy.deepcopy(child))
    for child in list(d_csld):
        if child is not d_tree:
            d_csld.remove(child)
    pos = 0
    for child in s_csld:
        if child.tag == f"{{{p}}}spTree":
            pos = list(d_csld).index(d_tree) + 1
            continue
        d_csld.insert(pos, copy.deepcopy(child))
        pos += 1
    for k, v in s_csld.attrib.items():
        d_csld.set(k, v)
    for child in list(dst_el):
        if child is not d_csld:
            dst_el.remove(child)
    for child in src_el:
        if child.tag != f"{{{p}}}cSld":
            dst_el.append(copy.deepcopy(child))
    for k, v in src_el.attrib.items():
        dst_el.set(k, v)

    to_remove = []
    for el in list(dst_el.iter()):
        for attr, val in list(el.attrib.items()):
            if not attr.startswith("{%s}" % R_NS) or val not in mapping:
                continue
            if mapping[val] is None:
                if drop_internal_links and str(el.tag).endswith(("hlinkClick", "hlinkHover")):
                    to_remove.append(el)
                continue
            el.set(attr, mapping[val])
    for el in to_remove:
        if el.getparent() is not None:
            el.getparent().remove(el)
    if foreign:
        materialize_placeholders(src_slide, new)
        fix_theme_colors(src_slide, new)
    return new


def theme_colors(slide):
    """{scheme name: RGB} for the slide's master, with the clrMap applied
    (tx1, bg1, tx2, bg2 resolved to dk1/lt1/dk2/lt2)."""
    master = slide.slide_layout.slide_master
    theme_part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    a = NS["a"]
    scheme = root.find(f".//{{{a}}}clrScheme")
    base = {}
    for child in scheme:
        name = child.tag.split("}")[1]
        c = child[0]
        val = c.get("lastClr") if c.tag.endswith("sysClr") else c.get("val")
        base[name] = (val or "000000").upper()
    cmap = master._element.find(f"{{{NS['p']}}}clrMap")
    out = dict(base)
    if cmap is not None:
        for k, v in cmap.attrib.items():
            if v in base:
                out[k] = base[v]
    return out


def default_text_rgb(slide, colors=None):
    """The colour plain text gets in the source template (master otherStyle,
    then the presentation default text style), as RGB."""
    a, p = NS["a"], NS["p"]
    colors = colors or theme_colors(slide)
    master = slide.slide_layout.slide_master._element
    pres = slide.part.package.presentation_part._element
    for root in (master.find(f".//{{{p}}}otherStyle"), pres.find(f"{{{p}}}defaultTextStyle")):
        if root is None:
            continue
        d = root.find(f"{{{a}}}lvl1pPr/{{{a}}}defRPr/{{{a}}}solidFill")
        if d is None or not len(d):
            continue
        c = d[0]
        if c.tag.endswith("srgbClr"):
            return c.get("val").upper()
        if c.tag.endswith("schemeClr") and colors.get(c.get("val")):
            return colors[c.get("val")]
    return "061F32"


def fix_theme_colors(src_slide, new_slide):
    """Freeze theme colours to the source template's RGB values, so a slide
    copied into a deck with another theme keeps its look."""
    a, p = NS["a"], NS["p"]
    colors = theme_colors(src_slide)
    for el in list(new_slide._element.iter(f"{{{a}}}schemeClr")):
        rgb = colors.get(el.get("val"))
        if not rgb:
            continue
        new = etree.Element(f"{{{a}}}srgbClr", val=rgb)
        for child in el:
            new.append(child)
        el.getparent().replace(el, new)
    text_rgb = default_text_rgb(src_slide, colors)
    fills = (f"{{{a}}}solidFill", f"{{{a}}}gradFill", f"{{{a}}}noFill")
    for sp in new_slide._element.iter(f"{{{p}}}sp"):
        if sp.find(f"{{{p}}}style/{{{a}}}fontRef") is not None and len(sp.find(f"{{{p}}}style/{{{a}}}fontRef")):
            continue  # the shape style sets the text colour
        tx = sp.find(f"{{{p}}}txBody")
        if tx is None:
            continue
        lst = tx.find(f"{{{a}}}lstStyle")
        if lst is not None and any(e.tag in fills for e in lst.iter()):
            continue
        for r in tx.iter(f"{{{a}}}r"):
            rpr = r.find(f"{{{a}}}rPr")
            if rpr is None:
                rpr = etree.Element(f"{{{a}}}rPr")
                r.insert(0, rpr)
            ppr = r.getparent().find(f"{{{a}}}pPr")
            if any(c.tag in fills for c in rpr) or (ppr is not None and any(e.tag in fills for e in ppr.iter())):
                continue
            fill = etree.Element(f"{{{a}}}solidFill")
            etree.SubElement(fill, f"{{{a}}}srgbClr", val=text_rgb)
            pos = 0
            for i, c in enumerate(rpr):
                if c.tag.endswith(("}ln",)):
                    pos = i + 1
            rpr.insert(pos, fill)


def materialize_placeholders(src_slide, new_slide):
    """The target deck has no matching layout (a deck made from another
    template): turn placeholders into plain shapes that carry the position
    and text style they inherited from the source layout and master."""
    p, a = NS["p"], NS["a"]
    src_layout = src_slide.slide_layout
    master = src_layout.slide_master._element
    styles = {"title": master.find(f".//{{{p}}}titleStyle"), "body": master.find(f".//{{{p}}}bodyStyle")}
    lay_ph = {ph.placeholder_format.idx: ph for ph in src_layout.placeholders}
    for sp in list(new_slide.shapes._spTree.iter(f"{{{p}}}sp")):
        ph = sp.find(f"{{{p}}}nvSpPr/{{{p}}}nvPr/{{{p}}}ph")
        if ph is None:
            continue
        idx = int(ph.get("idx", "0"))
        kind = ph.get("type", "body")
        src = lay_ph.get(idx)
        spPr = sp.find(f"{{{p}}}spPr")
        if src is not None and spPr is not None and spPr.find(f"{{{a}}}xfrm") is None:
            xfrm = src._element.find(f"{{{p}}}spPr/{{{a}}}xfrm")
            if xfrm is not None:
                spPr.insert(0, copy.deepcopy(xfrm))
        if spPr is not None and spPr.find(f"{{{a}}}prstGeom") is None and spPr.find(f"{{{a}}}custGeom") is None:
            geom = etree.Element(f"{{{a}}}prstGeom", prst="rect")
            etree.SubElement(geom, f"{{{a}}}avLst")
            xf = spPr.find(f"{{{a}}}xfrm")
            spPr.insert(list(spPr).index(xf) + 1 if xf is not None else 0, geom)
        tx = sp.find(f"{{{p}}}txBody")
        if tx is not None:
            lvl1 = None
            if src is not None:
                lvl1 = src._element.find(f".//{{{a}}}lstStyle/{{{a}}}lvl1pPr")
            if lvl1 is None:
                st = styles["title"] if kind in ("title", "ctrTitle") else styles["body"]
                lvl1 = st.find(f"{{{a}}}lvl1pPr") if st is not None else None
            lst = tx.find(f"{{{a}}}lstStyle")
            if lst is None:
                lst = etree.Element(f"{{{a}}}lstStyle")
                tx.insert(1, lst)
            st = styles["title"] if kind in ("title", "ctrTitle") else styles["body"]
            master_lvl1 = st.find(f"{{{a}}}lvl1pPr") if st is not None else None
            if lvl1 is not None and lst.find(f"{{{a}}}lvl1pPr") is None:
                lv = copy.deepcopy(lvl1)
                d0 = lv.find(f"{{{a}}}defRPr")
                md = master_lvl1.find(f"{{{a}}}defRPr") if master_lvl1 is not None else None
                if d0 is not None and md is not None:
                    # inherit what the layout leaves to the master: colour, size, bold
                    if d0.find(f"{{{a}}}solidFill") is None and md.find(f"{{{a}}}solidFill") is not None:
                        d0.insert(0, copy.deepcopy(md.find(f"{{{a}}}solidFill")))
                    for attr in ("sz", "b"):
                        if attr not in d0.attrib and attr in md.attrib:
                            d0.set(attr, md.get(attr))
                for latin in lv.iter(f"{{{a}}}latin"):
                    if latin.get("typeface", "").startswith("+"):
                        latin.set("typeface", "Arial")
                d = lv.find(f"{{{a}}}defRPr")
                if d is not None and d.find(f"{{{a}}}latin") is None:
                    etree.SubElement(d, f"{{{a}}}latin", typeface="Arial")
                lst.append(lv)
            if src is not None:
                sb = src._element.find(f".//{{{a}}}bodyPr")
                db = tx.find(f"{{{a}}}bodyPr")
                if sb is not None and db is not None and not db.attrib:
                    for k2, v2 in sb.attrib.items():
                        db.set(k2, v2)
        ph.getparent().remove(ph)


# --------------------------------------------------------------------------
# Shapes and slots
# --------------------------------------------------------------------------

def iter_shapes(shapes):
    """Yield every shape, descending into groups."""
    for sh in shapes:
        yield sh
        try:
            is_group = sh.shape_type == 6
        except NotImplementedError:
            is_group = False
        if is_group:
            yield from iter_shapes(sh.shapes)


def slot_shapes(slide) -> dict:
    """Map slot name -> shape for every shape named slot:<name>."""
    out = {}
    for sh in iter_shapes(slide.shapes):
        if sh.name.startswith("slot:"):
            out[sh.name[5:]] = sh
    return out


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def placeholder(slide, idx):
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == idx:
            return sh
    return None


# --------------------------------------------------------------------------
# Text filling that keeps the pattern's formatting
# --------------------------------------------------------------------------

_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _runs_from_markup(text):
    """Split 'Plain **bold** plain' into [(text, bold)]."""
    parts, pos = [], 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], None))
        parts.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], None))
    return parts or [("", None)]


def _para_protos(txBody):
    """(pPr, rPr) prototypes from the first two paragraphs that have runs."""
    a = NS["a"]
    protos = []
    for p_el in txBody.findall(f"{{{a}}}p"):
        r = p_el.find(f"{{{a}}}r")
        if r is None or not (r.findtext(f"{{{a}}}t") or "").strip():
            continue
        protos.append((p_el.find(f"{{{a}}}pPr"), r.find(f"{{{a}}}rPr"), p_el.find(f"{{{a}}}endParaRPr")))
        if len(protos) == 2:
            break
    if not protos:
        p_el = txBody.find(f"{{{a}}}p")
        if p_el is not None:
            epr = p_el.find(f"{{{a}}}endParaRPr")
            rpr = None
            if epr is not None:
                rpr = copy.deepcopy(epr)
                rpr.tag = f"{{{a}}}rPr"
            protos.append((p_el.find(f"{{{a}}}pPr"), rpr, epr))
        else:
            protos.append((None, None, None))
    return protos


def _style_key(proto):
    ppr, rpr, _ = proto
    keys = []
    for el in (ppr, rpr):
        if el is None:
            keys.append(None)
        else:
            keys.append(tuple(sorted((k, v) for k, v in el.attrib.items() if k not in ("lang", "dirty", "err"))))
    return tuple(keys)


def set_text(shape, value):
    """Replace a shape's text, keeping the pattern's formatting.

    value: str (newlines start new paragraphs) or list of str.
    **bold** marks bold spans. A paragraph starting with "## " uses the
    style of the box's first paragraph (heading); other paragraphs use the
    second paragraph's style when the box had two styles.
    """
    if isinstance(value, (list, tuple)):
        paras = [str(v) for v in value]
    else:
        paras = str(value).split("\n")
    tf = shape.text_frame
    txBody = tf._txBody
    a = NS["a"]
    protos = _para_protos(txBody)
    distinct = len(protos) == 2 and _style_key(protos[0]) != _style_key(protos[1])
    explicit = any(p.startswith("## ") for p in paras)
    two_styles = len(protos) == 2 and (explicit or (distinct and len(paras) > 1))
    if two_styles and not explicit:
        paras = ["## " + paras[0]] + list(paras[1:])
    for p_el in txBody.findall(f"{{{a}}}p"):
        txBody.remove(p_el)
    for text in paras:
        if text.startswith("## "):
            ppr, rpr, epr = protos[0]
            text = text[3:]
        elif two_styles:
            ppr, rpr, epr = protos[1]
        else:
            ppr, rpr, epr = protos[0]
        p_el = etree.SubElement(txBody, f"{{{a}}}p")
        if ppr is not None:
            p_el.append(copy.deepcopy(ppr))
        if text:
            for li, line in enumerate(text.split("\v")):
                if li:
                    br = etree.SubElement(p_el, f"{{{a}}}br")
                    if rpr is not None:
                        br.append(copy.deepcopy(rpr))
                chunks = _runs_from_markup(line)
                marked = any(b for _, b in chunks)
                for chunk, bold in chunks:
                    if not chunk:
                        continue
                    r = etree.SubElement(p_el, f"{{{a}}}r")
                    new_rpr = copy.deepcopy(rpr) if rpr is not None else etree.Element(f"{{{a}}}rPr")
                    new_rpr.set("lang", new_rpr.get("lang", "en-US"))
                    if bold:
                        new_rpr.set("b", "1")
                    elif marked:
                        new_rpr.set("b", "0")  # the team marked what is bold
                    r.append(new_rpr)
                    t = etree.SubElement(r, f"{{{a}}}t")
                    t.text = chunk
        if epr is not None:
            p_el.append(copy.deepcopy(epr))
        elif rpr is not None:
            end = copy.deepcopy(rpr)
            end.tag = f"{{{a}}}endParaRPr"
            p_el.append(end)


def shape_text(shape) -> str:
    if shape.has_text_frame:
        return shape.text_frame.text
    if getattr(shape, "has_table", False) and shape.has_table:
        return "\n".join(c.text for row in shape.table.rows for c in row.cells)
    return ""


def slide_text(slide, include_notes=False) -> str:
    parts = [shape_text(sh) for sh in iter_shapes(slide.shapes)]
    if include_notes and slide.has_notes_slide:
        parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(p for p in parts if p)


def font_sizes(shape):
    """Explicit font sizes (pt) used in a shape's runs."""
    out = []
    if not shape.has_text_frame:
        return out
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                out.append(r.font.size.pt)
    return out


def first_font_pt(shape, default=12.0):
    sizes = font_sizes(shape)
    if sizes:
        return sizes[0]
    a = NS["a"]
    for tag in ("endParaRPr", "defRPr"):
        el = shape._element.find(f".//{{{a}}}{tag}")
        if el is not None and el.get("sz"):
            return int(el.get("sz")) / 100
    return default


def capacity_chars(width_emu, height_emu, font_pt, char_w=0.52, line_h=1.2, bold=False):
    """Rough number of characters that fit in a box at font_pt (after the
    default 0.1" inner margins). Bold and big text are wider."""
    if not width_emu or not height_emu or not font_pt:
        return 0
    if font_pt >= 20:
        char_w = max(char_w, 0.6)
    elif bold:
        char_w = max(char_w, 0.55)
    w_pt = max(1.0, Emu(width_emu).pt - 14.4)
    h_pt = max(font_pt * line_h, Emu(height_emu).pt - 7.2)
    per_line = max(1, int(w_pt / (font_pt * char_w)))
    lines = max(1, int(h_pt / (font_pt * line_h)))
    return per_line * lines


def estimate_lines(text, width_emu, font_pt, char_w=0.52):
    w_pt = Emu(width_emu).pt
    per_line = max(1, int(w_pt / (font_pt * char_w)))
    total = 0
    for para in text.split("\n"):
        total += max(1, -(-len(para) // per_line))
    return total


def file_sha(path) -> str:
    return hashlib.sha1(Path(path).read_bytes()).hexdigest()


def slide_xml(slide) -> bytes:
    return etree.tostring(slide._element, method="c14n")


# --------------------------------------------------------------------------
# Geometry
# --------------------------------------------------------------------------

def bbox(shape):
    return (shape.left or 0, shape.top or 0,
            (shape.left or 0) + (shape.width or 0), (shape.top or 0) + (shape.height or 0))


def union(boxes):
    boxes = list(boxes)
    return (min(b[0] for b in boxes), min(b[1] for b in boxes),
            max(b[2] for b in boxes), max(b[3] for b in boxes))


def intersects(a, b, tol=0):
    return not (a[2] + tol < b[0] or b[2] + tol < a[0] or a[3] + tol < b[1] or b[3] + tol < a[1])


def center(b):
    return ((b[0] + b[2]) / 2, (b[1] + b[3]) / 2)


def inside(pt, b, tol=0):
    return b[0] - tol <= pt[0] <= b[2] + tol and b[1] - tol <= pt[1] <= b[3] + tol


IN = 914400  # EMU per inch


def is_line(shape):
    return shape.shape_type == 9 or shape._element.tag.endswith("}cxnSp")


def ungroup(slide, group):
    """Move a group's children to the slide with absolute coordinates."""
    a = NS["a"]
    g = group._element
    xfrm = g.find(f"{{{NS['p']}}}grpSpPr/{{{a}}}xfrm")
    off, ext = xfrm.find(f"{{{a}}}off"), xfrm.find(f"{{{a}}}ext")
    choff, chext = xfrm.find(f"{{{a}}}chOff"), xfrm.find(f"{{{a}}}chExt")
    ox, oy = int(off.get("x")), int(off.get("y"))
    ex, ey = int(ext.get("cx")) or 1, int(ext.get("cy")) or 1
    cx, cy = int(choff.get("x")), int(choff.get("y"))
    cex, cey = int(chext.get("cx")) or 1, int(chext.get("cy")) or 1
    sx, sy = ex / cex, ey / cey
    parent = g.getparent()
    pos = parent.index(g)
    for child in list(g):
        if child.tag.endswith(("}nvGrpSpPr", "}grpSpPr")):
            continue
        cxfrm = child.find(f".//{{{a}}}xfrm")
        if cxfrm is not None:
            co, ce = cxfrm.find(f"{{{a}}}off"), cxfrm.find(f"{{{a}}}ext")
            if co is not None:
                co.set("x", str(int(ox + (int(co.get("x")) - cx) * sx)))
                co.set("y", str(int(oy + (int(co.get("y")) - cy) * sy)))
            if ce is not None:
                ce.set("cx", str(int(int(ce.get("cx")) * sx)))
                ce.set("cy", str(int(int(ce.get("cy")) * sy)))
        parent.insert(pos, child)
        pos += 1
    parent.remove(g)


def top_level_ancestor(slide, shape):
    """The top-level shape that contains shape (itself if top-level)."""
    el = shape._element
    tree = slide.shapes._spTree
    while el.getparent() is not None and el.getparent() is not tree:
        el = el.getparent()
    for sh in slide.shapes:
        if sh._element is el:
            return sh
    return shape


def delete_with_decor(slide, shape, protect=()):
    """Delete a slot shape plus small decoration sitting on it and lines that
    only connected it. Shapes named in protect are never deleted."""
    target = top_level_ancestor(slide, shape)
    if target._element is not shape._element:
        # a slot inside a group: remove just that shape, keep its siblings
        el = shape._element
        if el.getparent() is not None:
            el.getparent().remove(el)
        return
    tb = bbox(target)
    area = max(1, (tb[2] - tb[0]) * (tb[3] - tb[1]))
    tol = int(0.12 * IN)
    victims = [target]
    rest = [s for s in slide.shapes if s._element is not target._element]
    for s in rest:
        if s.name.startswith("slot:") or s.name in protect or s.is_placeholder:
            continue
        b = bbox(s)
        if is_line(s):
            continue
        a2 = (b[2] - b[0]) * (b[3] - b[1])
        if a2 < area * 0.6 and intersects(b, tb, tol) and inside(center(b), tb, tol * 2):
            victims.append(s)
    # lines: delete when they touch the target and at most one other survivor
    survivors = [s for s in rest if s not in victims and not is_line(s)]
    for s in rest:
        if not is_line(s) or s.name in protect:
            continue
        b = bbox(s)
        if not intersects(b, tb, int(0.05 * IN)):
            continue
        touching = sum(1 for o in survivors if intersects(b, bbox(o), int(0.05 * IN)))
        if touching <= 1:
            victims.append(s)
    for v in victims:
        el = v._element
        if el.getparent() is not None:
            el.getparent().remove(el)


def redistribute(slide, groups, keep, axis="x", mode="move", span=None):
    """Spread the first `keep` repeat groups evenly across `span` (lo, hi in
    EMU along axis; defaults to the groups' own extent).
    groups: list of lists of top-level shapes, one list per group.
    mode: "move" keeps sizes; "resize" scales positions and sizes."""
    groups = [g for g in groups[:keep] if g]
    if not groups:
        return
    lo = 0 if axis == "x" else 1
    hi = lo + 2
    boxes = [union(bbox(s) for s in g) for g in groups]
    full_lo, full_hi = span if span else (min(b[lo] for b in boxes), max(b[hi] for b in boxes))
    n = len(groups)
    if mode == "resize":
        gap = max(0, boxes[1][lo] - boxes[0][hi]) if n > 1 else 0
        width = (full_hi - full_lo - gap * (n - 1)) / n
        for k, g in enumerate(groups):
            b = boxes[k]
            gw = max(1, b[hi] - b[lo])
            new_lo = full_lo + k * (width + gap)
            for s in g:
                sb = bbox(s)
                rel = (sb[lo] - b[lo]) / gw
                size = (sb[hi] - sb[lo]) * width / gw
                if axis == "x":
                    s.left, s.width = int(new_lo + rel * width), max(1, int(size))
                else:
                    s.top, s.height = int(new_lo + rel * width), max(1, int(size))
    else:
        step = (full_hi - full_lo) / n
        for k, g in enumerate(groups):
            b = boxes[k]
            d = int(full_lo + step * (k + 0.5) - (b[lo] + b[hi]) / 2)
            for s in g:
                if axis == "x":
                    s.left = (s.left or 0) + d
                else:
                    s.top = (s.top or 0) + d


# --------------------------------------------------------------------------
# Pictures, photos, charts, tables, Harvey balls
# --------------------------------------------------------------------------

def _image_bytes(img):
    if isinstance(img, (bytes, bytearray)):
        return bytes(img)
    return Path(img).read_bytes()


def _crop_to_fill(img_bytes, box_w, box_h):
    """srcRect values (l, t, r, b in 1/1000 %) so the image fills the box."""
    from PIL import Image
    im = Image.open(io.BytesIO(img_bytes))
    iw, ih = im.size
    if not iw or not ih or not box_w or not box_h:
        return (0, 0, 0, 0)
    ir, br = iw / ih, box_w / box_h
    if ir > br:  # too wide: crop left/right
        keep = br / ir
        c = int((1 - keep) / 2 * 100000)
        return (c, 0, c, 0)
    keep = ir / br
    c = int((1 - keep) / 2 * 100000)
    return (0, c, 0, c)


def replace_picture(slide, pic, img):
    """Swap the image of a picture shape, cropping to fill its box."""
    data = _image_bytes(img)
    _, rId = slide.part.get_or_add_image_part(io.BytesIO(data))
    a = NS["a"]
    blip = pic._element.find(f".//{{{a}}}blip")
    blip.set(f"{{{R_NS}}}embed", rId)
    for ext in blip.findall(f"{{{a}}}extLst"):
        blip.remove(ext)
    blipfill = blip.getparent()
    for old in blipfill.findall(f"{{{a}}}srcRect"):
        blipfill.remove(old)
    l, t, r, b = _crop_to_fill(data, pic.width, pic.height)
    src = etree.Element(f"{{{a}}}srcRect")
    for k, v in zip("ltrb", (l, t, r, b)):
        if v:
            src.set(k, str(v))
    blipfill.insert(list(blipfill).index(blip) + 1, src)
    cnv = pic._element.find(f".//{{{NS['p']}}}cNvPr")
    if cnv is not None and "descr" in cnv.attrib:
        del cnv.attrib["descr"]


def set_photo_fill(slide, shape, img=None, color="D9D9D9"):
    """Fill an autoshape (e.g. an oval) with a photo, or a flat colour."""
    a = NS["a"]
    spPr = shape._element.spPr
    for tag in ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"):
        for el in spPr.findall(f"{{{a}}}{tag}"):
            spPr.remove(el)
    geom = spPr.find(f"{{{a}}}prstGeom")
    if geom is None:
        geom = spPr.find(f"{{{a}}}custGeom")
    pos = list(spPr).index(geom) + 1 if geom is not None else len(spPr)
    if img is None:
        fill = etree.Element(f"{{{a}}}solidFill")
        etree.SubElement(fill, f"{{{a}}}srgbClr", val=color)
    else:
        data = _image_bytes(img)
        _, rId = slide.part.get_or_add_image_part(io.BytesIO(data))
        fill = etree.Element(f"{{{a}}}blipFill", rotWithShape="1")
        etree.SubElement(fill, f"{{{a}}}blip", {f"{{{R_NS}}}embed": rId})
        l, t, r, b = _crop_to_fill(data, shape.width, shape.height)
        src = etree.SubElement(fill, f"{{{a}}}srcRect")
        for k, v in zip("ltrb", (l, t, r, b)):
            if v:
                src.set(k, str(v))
        st = etree.SubElement(fill, f"{{{a}}}stretch")
        etree.SubElement(st, f"{{{a}}}fillRect")
    spPr.insert(pos, fill)


def drop_external_chart_data(chart_part):
    """Remove an externalData link (e.g. to a file on someone's laptop)."""
    c = NS["c"]
    root = chart_part._element
    for ed in root.findall(f"{{{c}}}externalData"):
        rId = ed.get(f"{{{R_NS}}}id")
        root.remove(ed)
        if rId and rId in chart_part.rels:
            chart_part.drop_rel(rId)


def fill_chart(shape, data):
    """data: {categories: [...], series: {name: [values]}, number_format?}"""
    from pptx.chart.data import CategoryChartData
    cd = CategoryChartData(number_format=data.get("number_format", "General"))
    cd.categories = [str(c) for c in data["categories"]]
    for name, values in data["series"].items():
        cd.add_series(str(name), [None if v is None else float(v) for v in values])
    drop_external_chart_data(shape.chart.part)
    shape.chart.replace_data(cd)
    c = NS["c"]
    for ax in shape.chart._chartSpace.iter(f"{{{c}}}valAx"):
        sc = ax.find(f"{{{c}}}scaling")
        for tag in ("max", "min"):
            for el in sc.findall(f"{{{c}}}{tag}") if sc is not None else []:
                sc.remove(el)
        for tag in ("majorUnit", "minorUnit"):
            for el in ax.findall(f"{{{c}}}{tag}"):
                ax.remove(el)


def fill_table(shape, rows):
    """Write rows (list of lists of str) into a native table, cell by cell,
    keeping each cell's formatting. Extra rows/columns are ignored."""
    tbl = shape.table
    a = NS["a"]

    def has_run(cell):
        return any((t.text or "").strip() for t in cell._tc.iter(f"{{{a}}}t"))

    for i, row in enumerate(rows[:len(tbl.rows)]):
        donors = [tbl.cell(i, j) for j in range(len(tbl.columns)) if has_run(tbl.cell(i, j))]
        for j, val in enumerate(row[:len(tbl.columns)]):
            cell = tbl.cell(i, j)
            val = "" if val is None else val
            if val == "" and not cell.text.strip():
                continue  # keep empty cells exactly as designed
            if not has_run(cell) and donors:
                body = cell._tc.find(f"{{{a}}}txBody")
                for p_el in body.findall(f"{{{a}}}p"):
                    body.remove(p_el)
                body.append(copy.deepcopy(donors[0]._tc.find(f".//{{{a}}}p")))
            set_text(cell, val)


def set_harvey(pie, oval, level, on="061F32", off="D9D9D9"):
    """Harvey ball level 0..4 using the template's pie-over-oval pair."""
    a = NS["a"]
    level = max(0, min(4, int(level)))

    def paint(shape, rgb):
        spPr = shape._element.spPr
        for tag in ("noFill", "solidFill", "gradFill"):
            for el in spPr.findall(f"{{{a}}}{tag}"):
                spPr.remove(el)
        fill = etree.Element(f"{{{a}}}solidFill")
        etree.SubElement(fill, f"{{{a}}}srgbClr", val=rgb)
        geom = spPr.find(f"{{{a}}}prstGeom")
        spPr.insert(list(spPr).index(geom) + 1, fill)

    paint(oval, on if level == 4 else off)
    if level in (0, 4):
        el = pie._element
        if el.getparent() is not None:
            el.getparent().remove(el)
        return
    paint(pie, on)
    # the pie must sit on top of its grey circle
    o_el, p_el = oval._element, pie._element
    if o_el.getparent() is p_el.getparent():
        par = o_el.getparent()
        if par.index(p_el) < par.index(o_el):
            par.remove(p_el)
            par.insert(par.index(o_el) + 1, p_el)
    start = 16200000  # 12 o'clock
    end = (start + level * 5400000) % 21600000
    av = pie._element.find(f".//{{{a}}}prstGeom/{{{a}}}avLst")
    for gd in list(av):
        av.remove(gd)
    etree.SubElement(av, f"{{{a}}}gd", name="adj1", fmla=f"val {start}")
    etree.SubElement(av, f"{{{a}}}gd", name="adj2", fmla=f"val {end}")


def placeholder_png(label, w=1600, h=900):
    """Neutral grey image with a label, used where a team photo goes."""
    from PIL import Image, ImageDraw, ImageFont
    im = Image.new("RGB", (w, h), (217, 217, 217))
    d = ImageDraw.Draw(im)
    text = f"[Image: {label}]"
    size = max(14, min(h // 16, int(w / (len(text) * 0.62))))
    try:
        font = ImageFont.truetype("Arial.ttf", size)
    except OSError:
        font = ImageFont.load_default()
    tw = d.textlength(text, font=font)
    d.text(((w - tw) / 2, h / 2 - h // 20), text, fill=(89, 89, 89), font=font)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


PLACEHOLDER_DESCR = "case-deck:placeholder-image"


def mark_placeholder_image(pic):
    cnv = pic._element.find(f".//{{{NS['p']}}}cNvPr")
    cnv.set("descr", PLACEHOLDER_DESCR)


def is_placeholder_image(shape):
    cnv = shape._element.find(f".//{{{NS['p']}}}cNvPr")
    return cnv is not None and cnv.get("descr") == PLACEHOLDER_DESCR


def verify_package(path):
    """Problems with a saved .pptx as a list of strings (empty = fine):
    duplicate zip entries, slide parts nobody points to, reload errors."""
    import zipfile
    problems = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        dups = {n for n in names if names.count(n) > 1}
        if dups:
            problems.append(f"duplicate entries: {sorted(dups)[:5]}")
        slide_parts = [n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
    try:
        prs = open_deck(path)
    except Exception as e:  # noqa: BLE001
        return problems + [f"does not reload: {e}"]
    if len(slide_parts) != len(prs.slides):
        problems.append(f"{len(slide_parts)} slide parts for {len(prs.slides)} slides (orphans)")
    return problems


def drop_unused_rels(slide):
    """Remove relationships no element points to any more (old images after
    a picture swap, photos replaced by a fill), so their files leave the deck."""
    keep = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}
    used = {v for el in slide._element.iter() for k, v in el.attrib.items() if k.startswith("{%s}" % R_NS)}
    for rId, rel in list(slide.part.rels.items()):
        if rel.reltype in keep or rId in used:
            continue
        slide.part.drop_rel(rId)
