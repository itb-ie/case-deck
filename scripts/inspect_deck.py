#!/usr/bin/env python3
"""Describe a deck as JSON: per slide, its layout, pattern, slots, text,
numbers, fonts, colours and shape bounds. review.py builds on this.

Usage:
  python inspect_deck.py deck.pptx              # JSON to stdout
  python inspect_deck.py deck.pptx --slides 3,7 --brief   # readable list
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import deckkit as dk  # noqa: E402
from pptx.util import Emu  # noqa: E402

NUM_RE = re.compile(r"(?<![\w.])(?:[A-Z]{3}\s?|[$€£¥]\s?)?[-+~]?\d[\d.,]*\s?(?:%|x|bn|m|k|M|B|K|pp|yrs?|y)?(?![\w])")


def _colors(shape):
    out = set()
    for el in shape._element.iter("{%s}srgbClr" % dk.NS["a"]):
        out.add(el.get("val", "").upper())
    return sorted(out)


def _fonts(shape):
    out = set()
    for el in shape._element.iter():
        tag = str(el.tag)
        if tag.endswith(("}latin", "}cs", "}ea")) and el.get("typeface"):
            tf = el.get("typeface")
            if not tf.startswith("+"):
                out.add(tf)
    return sorted(out)


def describe_shape(sh, depth=0):
    d = {
        "name": sh.name,
        "type": str(sh.shape_type).split(".")[-1].split(" ")[0] if sh.shape_type else "UNKNOWN",
        "left": sh.left, "top": sh.top, "width": sh.width, "height": sh.height,
        "depth": depth,
    }
    if sh.is_placeholder:
        d["ph_idx"] = sh.placeholder_format.idx
    text = dk.shape_text(sh)
    if text:
        d["text"] = text
        d["font_pt"] = dk.font_sizes(sh)
        if sh.has_text_frame:
            d["est_font_pt"] = dk.first_font_pt(sh)
    d["colors"] = _colors(sh)
    d["fonts"] = _fonts(sh)
    if getattr(sh, "has_chart", False) and sh.has_chart:
        ch = sh.chart
        d["chart"] = {"type": str(ch.chart_type).split(".")[-1].split(" ")[0],
                      "categories": [str(c) for c in ch.plots[0].categories] if ch.plots else [],
                      "series": [s.name for s in ch.series]}
    if getattr(sh, "has_table", False) and sh.has_table:
        d["table"] = [[c.text for c in row.cells] for row in sh.table.rows]
    if sh.shape_type == 6:
        d["children"] = [describe_shape(c, depth + 1) for c in sh.shapes]
    return d


def describe(path, slides=None):
    prs = dk.open_deck(path)
    out = {"path": str(path), "slide_count": len(prs.slides),
           "width": prs.slide_width, "height": prs.slide_height, "slides": []}
    for i, s in enumerate(prs.slides, 1):
        if slides and i not in slides:
            continue
        shapes = [describe_shape(sh) for sh in s.shapes]
        title = ""
        ph = dk.placeholder(s, 0)
        if ph is not None and ph.has_text_frame:
            title = ph.text_frame.text
        tracker = dk.placeholder(s, 17)
        all_text = dk.slide_text(s)
        out["slides"].append({
            "number": i,
            "slide_id": s.slide_id,
            "layout": s.slide_layout.name,
            "hidden": s._element.get("show") == "0",
            "pattern": _pattern_tag(s),
            "title": title,
            "tracker": tracker.text_frame.text if tracker is not None and tracker.has_text_frame else "",
            "slots": sorted(dk.slot_shapes(s).keys()),
            "numbers": sorted(set(m.group(0).strip() for m in NUM_RE.finditer(all_text))),
            "notes": s.notes_slide.notes_text_frame.text if s.has_notes_slide else "",
            "text": all_text,
            "shapes": shapes,
        })
    return out


def _pattern_tag(slide):
    for sh in slide.shapes:
        if sh.name.startswith("pattern:"):
            return sh.name[8:]
    m = slide._element.find(".//{%s}cNvPr[@descr]" % dk.NS["p"])
    return ""


def brief(info):
    lines = [f"{info['path']}  ({info['slide_count']} slides)"]
    for s in info["slides"]:
        lines.append(f"\n== slide {s['number']}  layout={s['layout']}  pattern={s['pattern'] or '-'}"
                     f"{'  HIDDEN' if s['hidden'] else ''}")
        if s["title"]:
            lines.append(f"   title: {s['title'][:110]}")

        def walk(shapes):
            for d in shapes:
                ind = "   " + "  " * d["depth"]
                pos = f"({Emu(d['left'] or 0).inches:.2f},{Emu(d['top'] or 0).inches:.2f} {Emu(d['width'] or 0).inches:.2f}x{Emu(d['height'] or 0).inches:.2f})"
                txt = d.get("text", "").replace("\n", " / ")[:70]
                extra = ""
                if "ph_idx" in d:
                    extra += f" ph{d['ph_idx']}"
                if d.get("font_pt"):
                    extra += f" {sorted(set(d['font_pt']))}pt"
                if "chart" in d:
                    extra += f" CHART {d['chart']['type']}"
                lines.append(f"{ind}{d['type'][:10]:10} {d['name'][:28]:28} {pos}{extra}  {txt}")
                if "children" in d:
                    walk(d["children"])
        walk(s["shapes"])
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--slides", help="1-based, e.g. 2,5-7")
    ap.add_argument("--brief", action="store_true", help="readable outline instead of JSON")
    args = ap.parse_args()
    slides = None
    if args.slides:
        from render import parse_range
        slides = set(parse_range(args.slides, 10_000))
    info = describe(args.pptx, slides)
    print(brief(info) if args.brief else json.dumps(info, indent=1, default=str))


if __name__ == "__main__":
    main()
