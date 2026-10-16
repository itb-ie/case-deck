#!/usr/bin/env python3
"""One-time curation: build assets/patterns.pptx and assets/catalog.yaml.

Takes golden slides from the 2025 winning decks (the house style) and a few
template samples, cleans them (case content, empty boxes, think-cell markers,
links to files on someone's laptop), renames every fillable shape to
slot:<name>, puts neutral sample text in, and records each pattern in the
catalogue (purpose, slots, max characters, repeatable groups).

Run from anywhere:
  python curate.py --winning "<path to winning_cases>"
The winning decks are only read, never written.
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
import deckkit as dk  # noqa: E402
from pptx.util import Emu  # noqa: E402

TAJIN = "Mexico 24H Case.pptx"
TPL = "PowerPoint Template.pptx"

# --------------------------------------------------------------------------
# Pattern specs
#
# slots: name -> selector | dict(sel=..., type=text|image|photo|chart|table|
#        harvey|marker, sample=..., optional=bool, max=int)
# selectors: "ph:0" placeholder idx | "t:Text" exact-or-prefix text match |
#            "n:Name" or "n:Name#2" shape name (nth) | "i:5" top-level index
# delete: selectors removed during curation
# repeat: [dict(prefix, count, min, axis, mode)] slots named <prefix><k>_*
# --------------------------------------------------------------------------

T = "## "  # heading-style paragraph marker in samples

PATTERNS = [
    dict(
        id="title", source=(TAJIN, 1),
        purpose="Cover slide: full-width photo, deck title, team photos and names, client logo.",
        when="First slide of every deck.",
        slots={
            "title": dict(sel="t:Mejorando", sample="[Deck title: short and evocative]", max=45),
            "hero_image": dict(sel="n:Picture 6", type="image", sample="hero photo of the client"),
            "client_logo": dict(sel="n:Picture 5", type="image", sample="client logo", optional=True),
            "m1_photo": dict(sel="n:Oval 24", type="photo", optional=True),
            "m2_photo": dict(sel="n:Oval 25", type="photo", optional=True),
            "m3_photo": dict(sel="n:Oval 22", type="photo", optional=True),
            "m4_photo": dict(sel="n:Oval 26", type="photo", optional=True),
            "m1_name": dict(sel="pos:1.67,6.97", sample="[Name 1]"),
            "m2_name": dict(sel="pos:3.43,6.97", sample="[Name 2]"),
            "m3_name": dict(sel="pos:5.2,6.97", sample="[Name 3]"),
            "m4_name": dict(sel="pos:7.02,6.96", sample="[Name 4]"),
        },
        clear=["n:Google Shape;185;p1"],
        repeat=[dict(prefix="m", count=4, min=2, axis="x", mode="move")],
    ),
    dict(
        id="exec_summary", source=(TAJIN, 2),
        purpose="Answer-first overview: the central question, case questions as pillars, one recommendation per pillar, headline impact.",
        when="Slide 2, and again just before the closing slide.",
        slots={
            "title": dict(sel="ph:0", sample="An overview of our approach"),
            "tracker": dict(sel="ph:17", sample="Executive Summary"),
            "question": dict(sel="t:How can Tajín stay", sample="[The central question of the case, in one sentence]"),
            "c1_num": dict(sel="n:Oval 36", sample="1"),
            "c2_num": dict(sel="n:Oval 37", sample="2"),
            "c3_num": dict(sel="n:Oval 38", sample="3"),
            "c1_head": dict(sel="t:Sabores de México", sample="[Pillar 1 name]"),
            "c2_head": dict(sel="t:Alma de Jalisco", sample="[Pillar 2 name]"),
            "c3_head": dict(sel="t:Llegar al mundo", sample="[Pillar 3 name]"),
            "c1_question": dict(sel="n:Rectangle 17", sample="[Case question 1 as asked]"),
            "c2_question": dict(sel="n:Rectangle 21", sample="[Case question 2 as asked]"),
            "c3_question": dict(sel="n:Rectangle 23", sample="[Case question 3 as asked]"),
            "c1_reco": dict(sel="n:Rectangle 25", sample=["**[Recommendation 1 name]**", "[One line on what it does]"]),
            "c2_reco": dict(sel="n:Rectangle 27", sample=["**[Recommendation 2 name]**", "[One line on what it does]"]),
            "c3_reco": dict(sel="n:Rectangle 29", sample=["**[Recommendation 3 name]**", "[One line on what it does]"]),
            "impact": dict(sel="t:Impact: contribute", sample="Impact: [headline number, e.g. USD ~00m revenue by 2030]"),
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
        delete=["t:Findings and recommendations"],
        repeat=[dict(prefix="c", count=3, min=2, axis="x", mode="resize")],
    ),
    dict(
        id="section_divider", source=(TAJIN, 3),
        purpose="Section break with a full-bleed photo, the section name and who presents it.",
        when="Before Analysis, Strategy, Impact (each presenter hand-off).",
        slots={
            "section": dict(sel="i:8", sample="[Section]"),
            "presenter": dict(sel="i:10", sample="[Presenter name]", max=45),
            "image": dict(sel="n:Picture 2", type="image", sample="section photo"),
        },
        delete=["i:0", "i:1", "i:2", "i:3", "i:4", "i:5"],
        align_to={"section": "n:Rectangle 3", "presenter": "n:Rectangle 3"},
    ),
    dict(
        id="company_position", source=(TAJIN, 4),
        purpose="Where the company stands: 3 big numbers, what they mean, core competencies, and a growth-potential map.",
        when="First Analysis slide.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: what the company has and what it must decide now]"),
            "tracker": dict(sel="ph:17", sample="Analysis"),
            "left_head": dict(sel="t:While Tajín leads", sample="[While the company leads in ...]"),
            "right_head": dict(sel="t:…it must now", sample="[... it must now decide ...]"),
            "stat1_value": dict(sel="n:TextBox 76", sample="00+"),
            "stat2_value": dict(sel="n:TextBox 79", sample="00+"),
            "stat3_value": dict(sel="n:TextBox 78", sample="00+"),
            "stat1_label": dict(sel="n:TextBox 80", sample="[What number 1 counts]"),
            "stat2_label": dict(sel="n:TextBox 81", sample="[What number 2 counts]"),
            "stat3_label": dict(sel="n:TextBox 82", sample="[What number 3 counts]"),
            "strength1": dict(sel="n:TextBox 50", sample="[Strength 1 and what it allows]"),
            "strength2": dict(sel="n:TextBox 55", sample="[Strength 2 and what it allows]"),
            "strength3": dict(sel="n:TextBox 56", sample="[Strength 3 and what it allows]"),
            "comp_head": dict(sel="t:Core Competencies", sample="Core Competencies"),
            "comp1": dict(sel="t:Cultural Storytelling", sample="[Competency 1]"),
            "comp2": dict(sel="t:Agile Product", sample="[Competency 2]"),
            "comp3": dict(sel="t:Authentic Mexican Flavor", sample="[Competency 3]"),
            "map_head": dict(sel="t:Growth Potential", sample="[Y axis: e.g. Growth potential]"),
            "map_x": dict(sel="t:Cultural & Brand", sample="[X axis label]"),
            "map1": dict(sel="n:TextBox 109", sample="[Option A]", optional=True),
            "map2": dict(sel="n:TextBox 111", sample="[Option B]", optional=True),
            "map3": dict(sel="n:TextBox 6", sample="[Option C]", optional=True),
            "map4": dict(sel="n:TextBox 12", sample="[Option D]", optional=True),
            "map5": dict(sel="n:TextBox 17", sample="[Option E]", optional=True),
            "map6": dict(sel="n:TextBox 110", sample="[Option F]", optional=True),
            "source": dict(sel="t:Source:", sample="Source: [Source]"),
        },
        delete=["n:TextBox 40"],
    ),
    dict(
        id="trends", source=(TAJIN, 5),
        purpose="Market or customer trends: 2-3 columns, each with a photo, 2 sourced numbers and a one-line takeaway.",
        when="Analysis section, after the company slide.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: what customers now expect]"),
            "tracker": dict(sel="ph:17", sample="Analysis"),
            "t1_image": dict(sel="n:Picture 6", type="image", sample="trend 1 photo"),
            "t1_head": dict(sel="t:Effortless Access", sample="[Trend 1]"),
            "t1_stat1": dict(sel="t:+20%", sample="00%"),
            "t1_label1": dict(sel="t:Increase in AOV", sample="[What it measures]"),
            "t1_stat2": dict(sel="t:2x", sample="0x"),
            "t1_label2": dict(sel="t:More likely to buy", sample="[What it measures]"),
            "t1_takeaway": dict(sel="t:Convenience drives", sample="[So what for the client]"),
            "t2_image": dict(sel="n:Picture 2", type="image", sample="trend 2 photo"),
            "t2_head": dict(sel="t:Experience & Identity", sample="[Trend 2]"),
            "t2_stat1": dict(sel="t:61%", sample="00%"),
            "t2_label1": dict(sel="t:Reward brands", sample="[What it measures]"),
            "t2_stat2": dict(sel="t:64%", sample="00%"),
            "t2_label2": dict(sel="t:Prefer experiences", sample="[What it measures]"),
            "t2_takeaway": dict(sel="t:Experience & personalization", sample="[So what for the client]"),
            "t3_image": dict(sel="n:Picture 8", type="image", sample="trend 3 photo"),
            "t3_head": dict(sel="t:Authenticity & Purpose", sample="[Trend 3]"),
            "t3_stat1": dict(sel="t:75%", sample="00%"),
            "t3_label1": dict(sel="t:Prefer purpose-led", sample="[What it measures]"),
            "t3_stat2": dict(sel="t:>55%", sample="00%"),
            "t3_label2": dict(sel="t:Would pay extra", sample="[What it measures]"),
            "t3_takeaway": dict(sel="t:Authenticity and sustainable", sample="[So what for the client]"),
            "source": dict(sel="t:Source:", sample="Source: [Sources for every number]"),
        },
        delete=["n:AutoShape 2"],
        repeat=[dict(prefix="t", count=3, min=2, axis="x", mode="move")],
    ),
    dict(
        id="hook_image", source=(TAJIN, 6),
        purpose="The hook: one full-bleed photo, no text. The presenter tells the human story out loud.",
        when="After the company context, before the analysis gets deep.",
        slots={"image": dict(sel="n:Picture 5", type="image", sample="hook photo (a real person or moment)")},
        delete=["ph:0", "ph:1", "ph:17"],
    ),
    dict(
        id="matrix_takeaways", source=(TAJIN, 7),
        purpose="A 2x2 positioning map (e.g. products by profitability vs relevance) with key takeaways on the right.",
        when="Analysis: to show where to play.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: where the client should double down]"),
            "tracker": dict(sel="ph:17", sample="Analysis"),
            "chart_head": dict(sel="t:Global edible SKUs", sample="[What the map plots]"),
            "y_axis": dict(sel="t:Profitability &", sample="[Y axis]"),
            "x_axis": dict(sel="t:Gen Z & Alpha Brand", sample="[X axis]"),
            "takeaways_head": dict(sel="t:Key takeaways", sample="Key takeaways and action points"),
            "takeaways": dict(sel="t:Clásico fuels", sample=[
                "**[Takeaway 1]:** [what it means and what to do]", "",
                "**[Takeaway 2]:** [what it means and what to do]", "",
                "**[Takeaway 3]:** [what it means and what to do]"]),
            "item1": dict(sel="t:Clásico", sample="[Item 1]", optional=True),
            "item2": dict(sel="t:Salsa Yaya", sample="[Item 2]", optional=True),
            "item3": dict(sel="t:Habanero", sample="[Item 3]", optional=True),
            "item4": dict(sel="t:Chamoy", sample="[Item 4]", optional=True),
            "item5": dict(sel="t:Low-Sodium", sample="[Item 5]", optional=True),
            "item6": dict(sel="t:Minis", sample="[Item 6]", optional=True),
            "item7": dict(sel="t:Twist", sample="[Item 7]", optional=True),
            "item8": dict(sel="t:Snack Kits", sample="[Item 8]", optional=True),
            "source": dict(sel="t:Source:", sample="Source: [Source]"),
        },
        dots=["n:Picture 2", "n:Picture 65", "n:Picture 67", "n:Picture 6", "n:Picture 8",
              "n:Picture 12", "n:Picture 14", "n:Picture 16"],
    ),
    dict(
        id="three_cards", source=(TAJIN, 8),
        purpose="2-3 photo cards, each with a bold heading and a short paragraph (e.g. strengths x opportunities).",
        when="Analysis synthesis, or any 'three things' message.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: where strengths meet opportunities]"),
            "subtitle": dict(sel="ph:1", sample="[Guiding question for this slide]", optional=True),
            "tracker": dict(sel="ph:17", sample="Analysis"),
            "c1_image": dict(sel="n:Picture 2", type="image", sample="card 1 photo"),
            "c2_image": dict(sel="n:Picture 6", type="image", sample="card 2 photo"),
            "c3_image": dict(sel="n:Picture 13", type="image", sample="card 3 photo"),
            "c1_text": dict(sel="n:TextBox 8", sample=["## [Card 1 heading]", "", "[Two or three lines with the evidence and the **key phrase** in bold]"]),
            "c2_text": dict(sel="n:TextBox 9", sample=["## [Card 2 heading]", "", "[Two or three lines with the evidence and the **key phrase** in bold]"]),
            "c3_text": dict(sel="n:TextBox 10", sample=["## [Card 3 heading]", "", "[Two or three lines with the evidence and the **key phrase** in bold]"]),
            "source": dict(sel="t:Source:", sample="Source: [Source]"),
        },
        delete=["n:Google Shape;23;p40"],
        repeat=[dict(prefix="c", count=3, min=2, axis="x", mode="resize")],
    ),
    dict(
        id="takeaway_overlay", source=(TAJIN, 9),
        purpose="Internal: the dark 'Key Takeaway' band laid over a copy of a slide. Use the takeaway field on any slide instead.",
        when="Never directly; build.py adds it when a slide has `takeaway:`.",
        internal=True,
        keep_only=["n:Rectangle 3", "n:Group 5"],
        slots={
            "takeaway_head": dict(sel="t:Key Takeaway", sample="Key Takeaway"),
            "takeaway": dict(sel="t:Tajín can unite", sample="[The one sentence the judges should remember]"),
        },
    ),
    dict(
        id="prioritization", source=(TAJIN, 11),
        purpose="Impact x feasibility matrix: the criteria used, the options as labelled dots in clusters, a legend and the conclusion.",
        when="Start of Strategy: show which options you chose and which you rejected.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: how we selected the strategies]"),
            "tracker": dict(sel="ph:17", sample="Strategy"),
            "criterion1": dict(sel="t:Revenue", sample="[Criterion 1]"),
            "criterion2": dict(sel="t:Implementation feasibility", sample="[Criterion 2]"),
            "criterion3": dict(sel="t:Investment required", sample="[Criterion 3]"),
            "criterion4": dict(sel="t:Organizational capabilities", sample="[Criterion 4]"),
            "legend1": dict(sel="t:High impact and feasible", sample=["High impact and feasible to implement", "[What these options have in common]"]),
            "legend2": dict(sel="t:High impact but", sample="High impact but requires high investment"),
            "legend3": dict(sel="t:Easy to implement", sample="Easy to implement but small impact"),
            "legend4": dict(sel="t:Low impact and difficult", sample="Low impact and difficult to implement"),
            "conclusion": dict(sel="t:Win the future", sample="[The selection logic in one sentence]"),
            **{f"opt{i}": dict(sel=f"t:{t}", sample=f"[Option {i}]", optional=True) for i, t in enumerate([
                "Savory snacks", "Popup stand", "Golden Ticket", "Zero waste", "Sweet snacks",
                "Targeted e-commerce", "Flavor co-creation", "Promote YaYa", "Intl. flavors",
                "Refilling bottles", "Points system", "Gateway flavors", "Digital flavor",
                "Dispensers in", "Online delivery", "Bottle recycling", "Flavor co-creation"], 1)},
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
    ),
    dict(
        id="strategy_tree", source=(TAJIN, 12),
        purpose="'Our perspective': the central question split into 3 areas, each with sub-areas and the initiatives under them.",
        when="Strategy section, after prioritisation. Delete boxes you do not need.",
        slots={
            "title": dict(sel="ph:0", sample="Our perspective on the strategic approaches [client] should follow"),
            "tracker": dict(sel="ph:17", sample="Strategy"),
            "question": dict(sel="n:Rectangle 4", sample="[“The central question of the case”]"),
            "b1": dict(sel="t:Product", sample="[Area 1]"),
            "b2": dict(sel="t:Brand", sample="[Area 2]"),
            "b3": dict(sel="t:Channel", sample="[Area 3]"),
            "b1_l1": dict(sel="t:Nurture Yaya", sample="[Initiative]", optional=True),
            "b1_l2": dict(sel="t:Transform international", sample="[Initiative]", optional=True),
            "b1_l3": dict(sel="t:Create snacks", sample="[Initiative]", optional=True),
            "b2_m1": dict(sel="t:Sustainability", sample="[Sub-area]", optional=True),
            "b2_m1_l1": dict(sel="t:Sustainable stories", sample="[Initiative]", optional=True),
            "b2_m1_l2": dict(sel="t:Zero waste", sample="[Initiative]", optional=True),
            "b2_m2": dict(sel="t:Community", sample="[Sub-area]", optional=True),
            "b2_m2_l1": dict(sel="t:Flavor co-creation", sample="[Initiative]", optional=True),
            "b3_m1": dict(sel="t:Online", sample="[Sub-area]", optional=True),
            "b3_m1_l1": dict(sel="t:Targeted", sample="[Initiative]", optional=True),
            "b3_m2": dict(sel="t:Experience", sample="[Sub-area]", optional=True),
            "b3_m2_l1": dict(sel="t:Golden ticket", sample="[Initiative]", optional=True),
            "b3_m2_l2": dict(sel="t:Pop-up stand", sample="[Initiative]", optional=True),
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
        delete=["n:AutoShape 15"],
    ),
    dict(
        id="issue_tree", source=(TPL, 174),
        purpose="Regular issue tree: main question, 2 drivers, 2 secondary drivers each, 2 underlying drivers each.",
        when="Analysis or appendix, when the structure is symmetric.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title for the issue tree]"),
            "tracker": dict(sel="ph:17", sample="Analysis"),
            "root": dict(sel="t:[Main issue]", sample="[Main question]"),
            "d1": dict(sel="n:Rectangle 5", sample="[Driver 1]"),
            "d2": dict(sel="n:Rectangle 6", sample="[Driver 2]"),
            "d1_s1": dict(sel="n:Rectangle 10", sample="[Secondary driver]", optional=True),
            "d1_s2": dict(sel="n:Rectangle 28", sample="[Secondary driver]", optional=True),
            "d2_s1": dict(sel="n:Rectangle 33", sample="[Secondary driver]", optional=True),
            "d2_s2": dict(sel="n:Rectangle 13", sample="[Secondary driver]", optional=True),
            "d1_s1_u1": dict(sel="n:Rectangle 15", sample="[Underlying driver]", optional=True),
            "d1_s1_u2": dict(sel="n:Rectangle 23", sample="[Underlying driver]", optional=True),
            "d1_s2_u1": dict(sel="n:Rectangle 29", sample="[Underlying driver]", optional=True),
            "d1_s2_u2": dict(sel="n:Rectangle 30", sample="[Underlying driver]", optional=True),
            "d2_s1_u1": dict(sel="n:Rectangle 34", sample="[Underlying driver]", optional=True),
            "d2_s1_u2": dict(sel="n:Rectangle 35", sample="[Underlying driver]", optional=True),
            "d2_s2_u1": dict(sel="n:Rectangle 26", sample="[Underlying driver]", optional=True),
            "d2_s2_u2": dict(sel="n:Rectangle 27", sample="[Underlying driver]", optional=True),
        },
        add_source=True,
    ),
    dict(
        id="framework", source=(TAJIN, 13),
        purpose="The strategy framework (acronym): 2-3 photo cards with number, pillar name and the 'How do we...?' question.",
        when="Strategy section. Also used as the pillar tracker (set `current:`).",
        slots={
            "title": dict(sel="ph:0", sample="A [ACRONYM] strategy to guide [client]"),
            "subtitle": dict(sel="t:Building Tajín", sample="[One line on what the framework achieves]"),
            "tracker": dict(sel="ph:17", sample="Strategy"),
            "p1_image": dict(sel="n:Picture 2", type="image", sample="pillar 1 photo"),
            "p2_image": dict(sel="n:Picture 4", type="image", sample="pillar 2 photo"),
            "p3_image": dict(sel="n:Picture 6", type="image", sample="pillar 3 photo"),
            "p1_shade": dict(sel="n:Rectangle 4", type="marker"),
            "p2_shade": dict(sel="n:Rectangle 7", type="marker"),
            "p3_shade": dict(sel="n:Rectangle 8", type="marker"),
            "p1_num": dict(sel="t:01", sample="01"),
            "p2_num": dict(sel="t:02", sample="02"),
            "p3_num": dict(sel="t:03", sample="03"),
            "p1_text": dict(sel="n:Rectangle 12", sample=["## [Pillar 1]", "[How can we ...?]"]),
            "p2_text": dict(sel="n:Rectangle 13", sample=["## [Pillar 2]", "[How can we ...?]"]),
            "p3_text": dict(sel="n:Rectangle 14", sample=["## [Pillar 3]", "[How can we ...?]"]),
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
        repeat=[dict(prefix="p", count=3, min=2, axis="x", mode="resize")],
    ),
    dict(
        id="pillar_overview", source=(TAJIN, 15),
        purpose="A pillar's key strategies: 2-3 rows, each a branded initiative, what it does, and a success case with numbers.",
        when="First slide of each pillar block.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: what this pillar achieves through which initiatives]"),
            "tracker": dict(sel="ph:17", sample="[Pillar]"),
            "left_head": dict(sel="t:Key strategies", sample="Key strategies"),
            "right_head": dict(sel="t:Success case study", sample="Success case study"),
            "i1_num": dict(sel="n:Oval 8", sample="1"),
            "i2_num": dict(sel="n:Oval 9", sample="2"),
            "i3_num": dict(sel="n:Oval 10", sample="3"),
            "i1_name": dict(sel="n:Rectangle 3", sample=["[Verb + object]", "**“[Branded name]”**"]),
            "i2_name": dict(sel="n:Rectangle 2", sample=["[Verb + object]", "**“[Branded name]”**"]),
            "i3_name": dict(sel="n:Rectangle 13", sample=["[Verb + object]", "**“[Branded name]”**"]),
            "i1_desc": dict(sel="t:Strengthen premium", sample="[What the initiative does, for whom, with the **key phrase** in bold]"),
            "i2_desc": dict(sel="t:Adapt signature", sample="[What the initiative does, for whom, with the **key phrase** in bold]"),
            "i3_desc": dict(sel="t:Collaborate with food", sample="[What the initiative does, for whom, with the **key phrase** in bold]"),
            "i1_case": dict(sel="t:Anchored quality", sample="**[Company]** [did what, with what result]"),
            "i2_case": dict(sel="t:“Worldwide favorites”", sample="**[Company]** [did what, with what result]"),
            "i3_case": dict(sel="t:Flamin’ Hot", sample="**[Company]** [did what, with what result]"),
            "i1_logo": dict(sel="n:Picture 2", type="image", sample="case logo", optional=True),
            "i2_logo": dict(sel="n:Picture 2#2", type="image", sample="case logo", optional=True),
            "i3_logo": dict(sel="n:Picture 6", type="image", sample="case logo", optional=True),
            "i1_stat1": dict(sel="t:195+", sample="00+"),
            "i1_label1": dict(sel="t:Countries sold", sample="[What it counts]"),
            "i1_stat2": dict(sel="t:30%", sample="00%", optional=True),
            "i1_label2": dict(sel="t:YoY increase", sample="[What it counts]", optional=True),
            "i2_stat1": dict(sel="t:6.5%", sample="0.0%"),
            "i2_label1": dict(sel="t:Increase in global", sample="[What it counts]"),
            "i3_stat1": dict(sel="t:$3B", sample="$0B"),
            "i3_label1": dict(sel="t:Retail sales", sample="[What it counts]"),
            "i3_stat2": dict(sel="t:12%", sample="00%", optional=True),
            "i3_label2": dict(sel="t:CAGR for spicy", sample="[What it counts]", optional=True),
            "source": dict(sel="t:Source:", sample="Source: [Sources of the cases]"),
        },
        repeat=[dict(prefix="i", count=3, min=2, axis="y", mode="none")],
    ),
    dict(
        id="initiative", source=(TAJIN, 18),
        purpose="One initiative in detail: breakdown in 3 points, partners, 3 KPIs with a target, and an example image.",
        when="Inside each pillar block, one per initiative.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: verb + initiative + result]"),
            "tracker": dict(sel="ph:17", sample="[Pillar]"),
            "breakdown_head": dict(sel="t:Breakdown of", sample="Breakdown of “[Branded name]”"),
            "b1_text": dict(sel="t:Develop co-branded", sample="[Step or feature 1, with the **key phrase** in bold]"),
            "b2_text": dict(sel="t:Synchronize drop", sample="[Step or feature 2, with the **key phrase** in bold]"),
            "b3_text": dict(sel="t:Target Gen Z", sample="[Who it targets, with the **key phrase** in bold]"),
            "partners_head": dict(sel="t:Potential partners", sample="Potential partners", optional=True),
            "partner1": dict(sel="n:Picture 4#2", type="image", sample="partner logo", optional=True),
            "partner2": dict(sel="n:Picture 4", type="image", sample="partner logo", optional=True),
            "partner3": dict(sel="n:Picture 16", type="image", sample="partner logo", optional=True),
            "partner4": dict(sel="n:Picture 2", type="image", sample="partner logo", optional=True),
            "kpis_head": dict(sel="t:KPIs", sample="KPIs"),
            "kpis": dict(sel="n:TextBox 1034", sample=[
                "## [KPI 1 name]", "[How it is measured] (target [number])", "",
                "## [KPI 2 name]", "[How it is measured] (target [number])", "",
                "## [KPI 3 name]", "[How it is measured] (target [number])"]),
            "example_head": dict(sel="t:Example", sample="Example"),
            "example_image": dict(sel="n:Picture 2#2", type="image", sample="example / mock-up"),
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
    ),
    dict(
        id="initiative_steps", source=(TAJIN, 17),
        purpose="An initiative as 3 steps (e.g. Discover, Prototype, Launch), with example images, a target-audience number and 2 KPIs.",
        when="Inside a pillar block, when the initiative is a process.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: verb + initiative + result]"),
            "subtitle": dict(sel="t:Keep the unmistakable", sample="[One line on the idea behind it]"),
            "tracker": dict(sel="ph:17", sample="[Pillar]"),
            "breakdown_head": dict(sel="t:Breakdown of", sample="Breakdown of “[Branded name]”"),
            "s1_name": dict(sel="n:Rectangle 24", sample=["[Step 1]", "[& verb]"]),
            "s2_name": dict(sel="n:Rectangle 22", sample=["[Step 2]", "[& verb]"]),
            "s3_name": dict(sel="n:Rectangle 26", sample=["[Step 3]", "[& verb]"]),
            "s1_text": dict(sel="n:TextBox 25", sample=["[Action]", "[Action]", "[Action]"]),
            "s2_text": dict(sel="n:TextBox 23", sample=["[Action]", "[Action]", "[Action]"]),
            "s3_text": dict(sel="n:TextBox 27", sample=["[Action]", "[Action]"]),
            "example_head": dict(sel="t:Example", sample="Example"),
            "ex1_image": dict(sel="n:Picture 4", type="image", sample="example 1"),
            "ex2_image": dict(sel="n:Picture 4#2", type="image", sample="example 2", optional=True),
            "ex3_image": dict(sel="n:Picture 4#3", type="image", sample="example 3", optional=True),
            "target_head": dict(sel="t:Target Audience", sample="Target Audience"),
            "target_stat": dict(sel="t:48%", sample="00%"),
            "target_label": dict(sel="n:TextBox 10", sample=["", "", "[Who, and what the number says about them]"]),
            "kpis_head": dict(sel="t:KPIs", sample="KPIs"),
            "kpi1": dict(sel="n:TextBox 39", sample=["## [KPI 1 name]", "[How it is measured]"]),
            "kpi2": dict(sel="n:TextBox 38", sample=["## [KPI 2 name]", "[How it is measured]"]),
            "source": dict(sel="t:Source:", sample="Source: [Source]"),
        },
        delete=["n:Picture 2", "n:Picture 2#2", "n:Picture 2#3", "n:Google Shape;23;p40"],
    ),
    dict(
        id="driver_tree", source=(TAJIN, 33),
        purpose="Revenue driver tree (e.g. penetration x frequency x price) linked to each area's initiatives by number.",
        when="First Impact slide: how each initiative moves a revenue driver.",
        slots={
            "title": dict(sel="ph:0", sample="To estimate financial impact, we link each strategy to the drivers of [client]'s revenue"),
            "tracker": dict(sel="ph:17", sample="Impact"),
            "left_head": dict(sel="t:We identify which", sample="We identify which factors impact revenue…"),
            "right_head": dict(sel="t:…and directly link", sample="…and directly link them to our strategies"),
            "root": dict(sel="t:Annual Revenue", sample="[Annual revenue]"),
            "d1": dict(sel="t:Penetration", sample="[Driver 1]"),
            "d2": dict(sel="t:Frequency", sample="[Driver 2]"),
            "d3": dict(sel="t:Average price", sample="[Driver 3]"),
            "op1": dict(sel="n:TextBox 31", sample="X"),
            "op2": dict(sel="n:TextBox 32", sample="X"),
            "a1": dict(sel="n:Rectangle 5", sample="[Area 1]"),
            "a2": dict(sel="n:Rectangle 87", sample="[Area 2]"),
            "a3": dict(sel="n:Rectangle 97", sample="[Area 3]"),
            "a1_s1": dict(sel="n:Rectangle 25", sample="[Initiative]"),
            "a1_s2": dict(sel="n:Rectangle 22", sample="[Initiative]", optional=True),
            "a1_s3": dict(sel="n:Rectangle 28", sample="[Initiative]", optional=True),
            "a2_s1": dict(sel="n:Rectangle 88", sample="[Initiative]"),
            "a2_s2": dict(sel="n:Rectangle 92", sample="[Initiative]", optional=True),
            "a2_s3": dict(sel="n:Rectangle 89", sample="[Initiative]", optional=True),
            "a3_s1": dict(sel="n:Rectangle 98", sample="[Initiative]"),
            "a3_s2": dict(sel="n:Rectangle 102", sample="[Initiative]", optional=True),
            "a3_s3": dict(sel="n:Rectangle 99", sample="[Initiative]", optional=True),
            "a1_s1_driver": dict(sel="n:Oval 60", sample="1"),
            "a1_s2_driver": dict(sel="n:Oval 84", sample="1", optional=True),
            "a1_s3_driver": dict(sel="n:Oval 86", sample="1", optional=True),
            "a2_s1_driver": dict(sel="n:Oval 94", sample="1"),
            "a2_s2_driver": dict(sel="n:Oval 95", sample="1", optional=True),
            "a2_s3_driver": dict(sel="n:Oval 96", sample="1", optional=True),
            "a3_s1_driver": dict(sel="n:Oval 104", sample="1"),
            "a3_s2_driver": dict(sel="n:Oval 105", sample="1", optional=True),
            "a3_s3_driver": dict(sel="n:Oval 106", sample="1", optional=True),
            "footnote": dict(sel="t:Footnote:", sample="Footnote: [What is included in each driver]", optional=True),
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
        repeat=[dict(prefix="a", count=3, min=2, axis="y", mode="none")],
    ),
    dict(
        id="financials", source=(TAJIN, 35),
        purpose="Financial impact: headline KPI pills (revenue, ROI, payback), key assumptions, driver increases and costs, and a base vs strategy chart.",
        when="Impact section, right after the driver tree.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: the strategy's financial return and why]"),
            "subtitle": dict(sel="ph:1", sample="[We forecast [client] reaching [number] in revenue by [year]]"),
            "tracker": dict(sel="ph:17", sample="Impact"),
            "kpi1_value": dict(sel="t:USD 135m", sample="USD 00m"),
            "kpi1_label": dict(sel="t:Revenue creation", sample="Revenue creation"),
            "kpi2_value": dict(sel="t:2.88x", sample="0.0x"),
            "kpi2_label": dict(sel="t:ROI", sample="ROI"),
            "kpi3_value": dict(sel="t:2y", sample="0y"),
            "kpi3_label": dict(sel="t:Payback", sample="Payback"),
            "as1_value": dict(sel="t:15%", sample="00%"),
            "as1_label": dict(sel="t:base growth", sample="[Assumption 1]"),
            "as2_value": dict(sel="t:45%", sample="00%", optional=True),
            "as2_label": dict(sel="t:penetration in Mexico", sample="[Assumption 2]", optional=True),
            "d1_name": dict(sel="n:Rectangle 46", sample="[Driver 1]"),
            "d2_name": dict(sel="n:Rectangle 48", sample="[Driver 2]"),
            "d3_name": dict(sel="n:Rectangle 47", sample="[Driver 3]"),
            "inc_head": dict(sel="t:Increase (6y)", sample="Increase (6y)"),
            "cost_head": dict(sel="t:Cost (USD, 6y)", sample="Cost (USD, 6y)"),
            "d1_inc": dict(sel="t:2%", sample="0%"),
            "d2_inc": dict(sel="pos:2.75,5.62", sample="0%"),
            "d3_inc": dict(sel="t:10%", sample="0%"),
            "d1_cost": dict(sel="t:~26.5m", sample="~0m"),
            "d2_cost": dict(sel="t:~13m", sample="~0m"),
            "d3_cost": dict(sel="t:~7.5m", sample="~0m"),
            "total_inc": dict(sel="t:~26%", sample="~0%"),
            "total_cost": dict(sel="t:~47m", sample="~0m"),
            "chart_head": dict(sel="t:Implementing our strategies", sample="[What the chart shows, with the main driver]"),
            "chart": dict(sel="n:Chart 14", type="chart"),
            "contrib1": dict(sel="t:8%", sample="0%"),
            "contrib2": dict(sel="t:12%", sample="0%"),
            "contrib3": dict(sel="n:Rectangle 8", sample="0%"),
            "contrib4": dict(sel="n:Rectangle 9", sample="0%"),
            "contrib5": dict(sel="t:17%", sample="0%"),
            "contrib6": dict(sel="n:Rectangle 12", sample="0%"),
            "growth_label": dict(sel="t:~19%", sample="~0%"),
            "footnote": dict(sel="t:Footnote:", sample="Footnote: ROI = total revenue from strategies / total cost", optional=True),
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
        chart_sample=dict(
            categories=["2024", "2025E", "2026E", "2027E", "2028E", "2029E", "2030E"],
            series={"Base growth": [300, 320, 340, 360, 380, 400, 420],
                    "Penetration": [0, 10, 20, 30, 40, 50, 60],
                    "Frequency": [0, 8, 16, 24, 32, 40, 48],
                    "Avg. price": [0, 4, 8, 12, 16, 20, 24]}),
    ),
    dict(
        id="chart_takeaways", source=(TPL, 11),
        purpose="A native column chart with a description on top and key takeaways on the right.",
        when="Analysis or appendix, any time a number series makes the point.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: what the chart proves]"),
            "subtitle": dict(sel="ph:1", sample="[Subtitle]", optional=True),
            "tracker": dict(sel="ph:17", sample="Analysis"),
            "chart_head": dict(sel="t:[Description]", sample="[What the chart shows]"),
            "unit": dict(sel="n:TextBox 7", sample="[Data], [Unit]"),
            "chart": dict(sel="n:Chart 16", type="chart"),
            "takeaways_head": dict(sel="t:[Key takeaways", sample="Key takeaways"),
            "takeaways": dict(sel="n:TextBox 17", sample=["[Takeaway 1]", "", "[Takeaway 2]", "", "[Takeaway 3]"]),
            "footnote": dict(sel="t:Footnote:", sample="Footnote: [Note]", optional=True),
            "source": dict(sel="t:Source:", sample="Source: [Source]"),
        },
        delete=["n:TextBox 10"],
        chart_sample=dict(categories=["[A]", "[B]", "[C]", "[D]", "[E]"],
                          series={"[Series]": [50, 40, 30, 20, 10]}),
    ),
    dict(
        id="harvey_table", source=(TPL, 72),
        purpose="Comparison table with Harvey balls (0-4): options or competitors against 5 criteria, plus a takeaway bar.",
        when="Analysis (competitors) or Strategy (option evaluation); also good in the appendix.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: who wins on what]"),
            "subtitle": dict(sel="ph:1", sample="[Subtitle]", optional=True),
            "tracker": dict(sel="ph:17", sample="Analysis"),
            "table": dict(sel="n:Table 3", type="table"),
            "takeaway": dict(sel="t:[Insert main takeaway]", sample="[Main takeaway]"),
            "source": dict(sel="t:Source:", sample="Source: [Source]"),
        },
        harvey=True,
    ),
    dict(
        id="risk_matrix", source=(TAJIN, 36),
        purpose="Risk matrix: risks A-G placed by probability and impact, coloured by pillar, with a table of descriptions and top-risk stars.",
        when="Impact section, before mitigation.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: the main risks to address]"),
            "tracker": dict(sel="ph:17", sample="Impact"),
            "legend1": dict(sel="t:Sabores", sample="[Pillar 1]"),
            "legend2": dict(sel="t:Alma", sample="[Pillar 2]"),
            "legend3": dict(sel="t:Llegada", sample="[Pillar 3]", optional=True),
            **{f"r{k}_name": dict(sel=f"n:{n}", sample=f"[Risk {chr(64 + k)}]", optional=k > 3)
               for k, n in enumerate(["Rectangle 43", "Rectangle 41", "Rectangle 9", "Rectangle 45",
                                      "Rectangle 70", "Rectangle 69", "Rectangle 71"], 1)},
            **{f"r{k}_desc": dict(sel=f"n:{n}", sample="[What could go wrong and why it matters]", optional=k > 3)
               for k, n in enumerate(["Rectangle 10", "Rectangle 82", "Rectangle 84", "Rectangle 90",
                                      "Rectangle 87", "Rectangle 88", "Rectangle 89"], 1)},
            **{f"r{k}_letter": dict(sel=f"n:{n}", type="marker", optional=k > 3)
               for k, n in enumerate(["Oval 72", "Oval 74", "Oval 75", "Oval 91", "Oval 92",
                                      "Oval 93", "Oval 94"], 1)},
            **{f"r{k}_dot": dict(sel=f"n:{n}", type="marker", optional=k > 3)
               for k, n in enumerate(["Oval 78", "Oval 77", "Oval 76", "Oval 95", "Oval 81",
                                      "Oval 80", "Oval 79"], 1)},
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
        stars=dict(existing=["n:5-point Star 30", "n:5-point Star 3", "n:5-point Star 2"], rows="r{k}_name"),
    ),
    dict(
        id="mitigation", source=(TAJIN, 37),
        purpose="Risk mitigation: for each top risk, the mitigation and the actionable steps.",
        when="Right after the risk matrix.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: how we reduce the main risks]"),
            "tracker": dict(sel="ph:17", sample="Impact"),
            "steps_head": dict(sel="t:Actionable steps", sample="Action steps"),
            "m1_letter": dict(sel="n:Oval 5", sample="A"),
            "m2_letter": dict(sel="n:Oval 8", sample="B"),
            "m3_letter": dict(sel="n:Oval 10", sample="C"),
            "m1_name": dict(sel="t:Identity dilution", sample="[Risk]"),
            "m2_name": dict(sel="t:Greenwashing", sample="[Risk]"),
            "m3_name": dict(sel="t:Community fatigue", sample="[Risk]"),
            "m1_mitigation": dict(sel="n:TextBox 46", sample="[Mitigation, with the **key phrase** in bold]"),
            "m2_mitigation": dict(sel="n:TextBox 49", sample="[Mitigation, with the **key phrase** in bold]"),
            "m3_mitigation": dict(sel="n:TextBox 50", sample="[Mitigation, with the **key phrase** in bold]"),
            "m1_steps": dict(sel="n:TextBox 3", sample=["[Step]", "[Step]", "[Step]"]),
            "m2_steps": dict(sel="n:TextBox 48", sample=["[Step]", "[Step]", "[Step]"]),
            "m3_steps": dict(sel="n:TextBox 23", sample=["[Step]", "[Step]", "[Step]"]),
            "footnote": dict(sel="t:Footnote:", sample="Footnote: Illustrative; not exhaustive", optional=True),
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
        repeat=[dict(prefix="m", count=3, min=2, axis="y", mode="none")],
    ),
    dict(
        id="roadmap", source=(TAJIN, 38),
        purpose="Implementation roadmap: 3-5 phases as chevrons with timing, what happens in each, and a rising phase line.",
        when="Impact section, last content slide before the recap.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title: pilot first, then scale]"),
            "tracker": dict(sel="ph:17", sample="Impact"),
            **{f"ph{k}_name": dict(sel=f"n:{n}", sample=f"[Phase {k} name]")
               for k, n in enumerate(["Chevron 167", "Chevron 10", "Chevron 14", "Chevron 17", "Chevron 103"], 1)},
            **{f"ph{k}_time": dict(sel=f"n:{n}", sample="Months 0-0")
               for k, n in enumerate(["TextBox 47", "TextBox 48", "TextBox 49", "TextBox 50", "TextBox 51"], 1)},
            **{f"ph{k}_text": dict(sel=f"n:{n}", sample=["[What happens: pilots, cities, numbers]", "", "[Milestone]"])
               for k, n in enumerate(["Rectangle 34", "Rectangle 36", "Rectangle 43", "Rectangle 44", "Rectangle 35"], 1)},
            **{f"ph{k}_label": dict(sel=f"t:Phase {k}", sample=f"Phase {k}")
               for k in range(1, 6)},
            "footnote": dict(sel="t:Footnote:", sample="Footnote: Illustrative; not exhaustive", optional=True),
            "source": dict(sel="t:Source:", sample="Source: Team Analysis"),
        },
        ungroup=["n:Group 52"],
        repeat=[dict(prefix="ph", count=5, min=3, axis="x", mode="resize")],
    ),
    dict(
        id="appendix_divider", source=(TAJIN, 42),
        purpose="Dark divider that opens the appendix.",
        when="After the closing slide, first slide of the appendix.",
        slots={
            "section": dict(sel="t:Appendix", sample="Appendix"),
            "client_logo": dict(sel="n:Picture 2", type="image", sample="client logo", optional=True),
        },
    ),
    dict(
        id="appendix_index", source=(TAJIN, 43),
        purpose="Appendix index. navigate.py fills it with Slide Zooms to every backup slide; without navigation it lists them.",
        when="Second appendix slide.",
        slots={
            "title": dict(sel="ph:0", sample="Appendix. Index"),
            "tracker": dict(sel="ph:17", sample="Appendix"),
            "index_list": dict(sel="n:Text Placeholder 5", sample=["[Backup slide 1]", "[Backup slide 2]"], optional=True),
        },
    ),
    dict(
        id="text_content", source=(TAJIN, 43),
        purpose="Plain content slide: action title, optional subtitle, bullets. Use when no other pattern fits.",
        when="Anywhere; appendix text slides.",
        slots={
            "title": dict(sel="ph:0", sample="[Action title]"),
            "tracker": dict(sel="ph:17", sample="[Section]"),
            "body": dict(sel="n:Text Placeholder 5", sample=["[Point 1]", "[Point 2]", "[Point 3]"]),
        },
        add_source=True,
    ),
    dict(
        id="closing", source=(TAJIN, 41),
        purpose="Closing slide: hero photo, team photos and names, client logo.",
        when="Last slide before the appendix.",
        slots={
            "hero_image": dict(sel="n:Picture 6", type="image", sample="closing photo"),
            "client_logo": dict(sel="n:Picture 5", type="image", sample="client logo", optional=True),
            "m1_photo": dict(sel="n:Oval 24", type="photo", optional=True),
            "m2_photo": dict(sel="n:Oval 25", type="photo", optional=True),
            "m3_photo": dict(sel="n:Oval 22", type="photo", optional=True),
            "m4_photo": dict(sel="n:Oval 26", type="photo", optional=True),
            "m1_name": dict(sel="pos:1.67,6.97", sample="[Name 1]"),
            "m2_name": dict(sel="pos:3.43,6.97", sample="[Name 2]"),
            "m3_name": dict(sel="pos:5.2,6.97", sample="[Name 3]"),
            "m4_name": dict(sel="pos:7.02,6.96", sample="[Name 4]"),
        },
        clear=["n:Google Shape;185;p1"],
        repeat=[dict(prefix="m", count=4, min=2, axis="x", mode="move")],
    ),
]


# --------------------------------------------------------------------------
# Selector resolution
# --------------------------------------------------------------------------

def _norm(s):
    return re.sub(r"\s+", " ", s or "").strip()


def find(slide, sel, taken=()):
    kind, _, arg = sel.partition(":")
    shapes = [s for s in dk.iter_shapes(slide.shapes) if s._element not in taken]
    if kind == "ph":
        return dk.placeholder(slide, int(arg))
    if kind == "i":
        return list(slide.shapes)[int(arg)]
    if kind == "n":
        name, _, nth = arg.partition("#")
        hits = [s for s in dk.iter_shapes(slide.shapes) if s.name == name]
        k = int(nth) - 1 if nth else 0
        return hits[k] if len(hits) > k else None
    if kind == "pos":
        x, y = (float(v) * dk.IN for v in arg.split(","))
        hits = [s for s in slide.shapes if s._element not in taken and s.shape_type != 6
                and dk.inside((x, y), dk.bbox(s))]
        hits.sort(key=lambda s: (s.width or 0) * (s.height or 0))
        return hits[0] if hits else None
    if kind == "t":
        target = _norm(arg)
        for s in shapes:
            if _norm(dk.shape_text(s)) == target:
                return s
        for s in shapes:
            if _norm(dk.shape_text(s)).startswith(target):
                return s
        return None
    raise ValueError(sel)


# --------------------------------------------------------------------------
# Cleaning helpers
# --------------------------------------------------------------------------

def is_junk(sh):
    """Empty or zero-size boxes, OLE markers (think-cell)."""
    el = sh._element
    if el.tag.endswith("}graphicFrame") and el.find(".//{%s}oleObj" % dk.NS["p"]) is not None:
        return True
    if sh.shape_type == 7:  # EMBEDDED_OLE_OBJECT
        return True
    if sh.has_text_frame and not sh.text_frame.text.strip() and not sh.is_placeholder:
        w, h = sh.width or 0, sh.height or 0
        a = dk.NS["a"]
        sp = el.find(f"{{{dk.NS['p']}}}spPr")
        has_fill = sp is not None and sp.find(f"{{{a}}}solidFill") is not None
        has_line = sp is not None and sp.find(f"{{{a}}}ln/{{{a}}}solidFill") is not None
        if w < 0.1 * dk.IN or h < 0.1 * dk.IN or not (has_fill or has_line):
            if sh.shape_type in (17, 1) and not has_fill:
                return True
    return False


def picture_to_dot(slide, pic):
    """Replace a small product picture by a cyan dot of similar size."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    size = int(min(pic.width, pic.height, 0.32 * dk.IN))
    cx, cy = pic.left + pic.width // 2, pic.top + pic.height // 2
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - size // 2, cy - size // 2, size, size)
    dot.fill.solid()
    dot.fill.fore_color.rgb = RGBColor(0x22, 0xA3, 0xDF)
    dot.line.fill.background()
    dot.name = "dot"
    pic._element.getparent().replace(pic._element, dot._element)


def copy_source_box(dst_slide, src_slide):
    import copy as _copy
    for sh in src_slide.shapes:
        if dk.shape_text(sh).strip().startswith("Source:"):
            el = _copy.deepcopy(sh._element)
            dst_slide.shapes._spTree.append(el)
            return dst_slide.shapes[-1]
    return None


def clear_authors(prs):
    cp = prs.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.title = ""


def default_sample(slot):
    if slot == "title":
        return "[Action title: the takeaway in one sentence, with a number]"
    if slot == "tracker":
        return "[Section]"
    if slot == "source":
        return "Source: [Source]"
    return "[" + slot.replace("_", " ").title() + "]"


# --------------------------------------------------------------------------
# Curation
# --------------------------------------------------------------------------

def curate(winning: Path, out_dir: Path, only=None):
    template = winning / TPL
    out_dir.mkdir(parents=True, exist_ok=True)
    tpl_copy = out_dir / "template.pptx"
    tpl = dk.open_deck(template)
    clear_authors(tpl)
    tpl.save(str(tpl_copy))  # same template, without the 2025 authors' names in the file properties

    dst = dk.open_deck(tpl_copy)
    for s in list(dst.slides):
        dk.drop_slide(dst, s)
    sources = {}
    catalog = {"version": 1, "slide_width": int(dst.slide_width), "slide_height": int(dst.slide_height),
               "patterns": {}}
    placeholder_cache = {}
    tajin_src = None

    for spec in PATTERNS:
        if only and spec["id"] not in only:
            continue
        fname, num = spec["source"]
        if fname not in sources:
            sources[fname] = dk.open_deck(winning / fname)
        src_prs = sources[fname]
        if fname == TAJIN:
            tajin_src = src_prs
        slide = dk.clone_slide(src_prs.slides[num - 1], dst)
        sid = spec["id"]

        # --- structural edits before naming
        for sel in spec.get("ungroup", []):
            g = find(slide, sel)
            if g is not None:
                dk.ungroup(slide, g)
        if "keep_only" in spec:
            keep = [find(slide, s) for s in spec["keep_only"]]
            keep_ids = {id(k._element) for k in keep if k is not None}
            for sh in list(slide.shapes):
                if id(sh._element) not in keep_ids:
                    dk.remove_shape(sh)

        # resolve all slot shapes first (selectors refer to original text)
        resolved, taken, orig_len = {}, set(), {}
        for slot, sd in spec["slots"].items():
            sd = sd if isinstance(sd, dict) else {"sel": sd}
            sh = find(slide, sd["sel"], taken)
            if sh is None:
                raise SystemExit(f"[{sid}] slot {slot}: selector {sd['sel']!r} matched nothing")
            if sh._element in taken:
                other = next(k for k, (v, _) in resolved.items() if v._element is sh._element)
                raise SystemExit(f"[{sid}] slots {other} and {slot} point at the same shape; make a selector more specific")
            taken.add(sh._element)
            resolved[slot] = (sh, sd)
            orig_len[slot] = len(_norm(dk.shape_text(sh)))
        to_delete = [find(slide, s) for s in spec.get("delete", [])]
        to_clear = [find(slide, s) for s in spec.get("clear", [])]
        dot_pics = [find(slide, s) for s in spec.get("dots", [])]
        stars = [find(slide, s) for s in spec.get("stars", {}).get("existing", [])]

        for sh in to_delete:
            if sh is not None and sh._element.getparent() is not None:
                dk.remove_shape(sh)
        for sh in to_clear:
            if sh is not None:
                dk.set_text(sh, "")
        for pic in dot_pics:
            if pic is not None:
                picture_to_dot(slide, pic)
        slot_ids = {id(sh._element) for sh, _ in resolved.values()}
        for sh in list(dk.iter_shapes(slide.shapes)):
            if id(sh._element) not in slot_ids and is_junk(sh) and sh._element.getparent() is not None:
                sh._element.getparent().remove(sh._element)

        # --- widen text boxes to a reference box and centre them
        for slot, sel in spec.get("align_to", {}).items():
            ref = find(slide, sel)
            sh = resolved[slot][0]
            if ref is not None:
                sh.left, sh.width = ref.left, ref.width
                for para in sh.text_frame.paragraphs:
                    from pptx.enum.text import PP_ALIGN
                    para.alignment = PP_ALIGN.CENTER

        # --- fill slots with samples, name them, collect catalogue entries
        entry_slots = {}
        for slot, (sh, sd) in resolved.items():
            stype = sd.get("type", "text")
            info = {"type": stype}
            if sd.get("optional"):
                info["optional"] = True
            if stype == "text":
                sample = sd.get("sample", default_sample(slot))
                if sh.has_text_frame:
                    dk.set_text(sh, sample)
                    pt = dk.first_font_pt(sh, 14.0)
                    if sh.is_placeholder and sh.placeholder_format.idx == 0:
                        info["max_chars"] = sd.get("max", 150)
                        pt = 25.0
                    else:
                        bold = any(r.font.bold for p_ in sh.text_frame.paragraphs for r in p_.runs)
                        cap = dk.capacity_chars(sh.width, sh.height, pt, char_w=0.5, bold=bold)
                        sample_len = len(_norm(" ".join(sample) if isinstance(sample, list) else sample))
                        info["max_chars"] = sd.get("max", max(cap, int(orig_len.get(slot, 0) * 1.15), sample_len, 8))
                    info["font_pt"] = pt
                info["sample"] = sample
            elif stype == "image":
                label = sd.get("sample", slot)
                key = (label, round((sh.width or 1) / (sh.height or 1), 2))
                if key not in placeholder_cache:
                    ratio = (sh.width or 1) / (sh.height or 1)
                    w = 1600 if ratio >= 1 else int(1600 * ratio)
                    h = int(w / ratio) if ratio >= 1 else 1600
                    placeholder_cache[key] = dk.placeholder_png(label, max(w, 200), max(h, 200))
                dk.replace_picture(slide, sh, placeholder_cache[key])
                dk.mark_placeholder_image(sh)
                info["label"] = label
                info["aspect"] = round((sh.width or 1) / (sh.height or 1), 3)
            elif stype == "photo":
                dk.set_photo_fill(slide, sh, None)
            elif stype == "chart":
                for dpt in list(sh.chart._chartSpace.iter("{%s}dPt" % dk.NS["c"])):
                    dpt.getparent().remove(dpt)  # no highlighted sample bar
                data = spec.get("chart_sample")
                if data:
                    dk.fill_chart(sh, data)
                    info["sample"] = data
                ch = sh.chart
                info["chart_type"] = str(ch.chart_type).split(".")[-1].split(" ")[0]
            elif stype == "table":
                tbl = sh.table
                info["rows"] = len(tbl.rows)
                info["cols"] = len(tbl.columns)
            sh.name = f"slot:{slot}"
            entry_slots[slot] = info

        # --- Harvey balls: name pie/oval pairs by row and column
        if spec.get("harvey"):
            tbl_shape = resolved["table"][0]
            tb = dk.bbox(tbl_shape)
            pies = [s for s in slide.shapes if s.name.startswith("Pie")]
            ovals = [s for s in slide.shapes if s.name.startswith("Oval")]
            cols = sorted({p.left for p in pies})
            rows = sorted({p.top for p in pies})
            for p in pies:
                r, c = rows.index(p.top) + 1, cols.index(p.left) + 1
                p.name = f"slot:h_r{r}_c{c}"
                match = [o for o in ovals if o.left == p.left and o.top == p.top]
                if match:
                    match[0].name = f"hb:r{r}_c{c}"
                dk.set_harvey(p, match[0], 2 if (r + c) % 2 else 1) if match else None
            for r in range(1, len(rows) + 1):
                for c in range(1, len(cols) + 1):
                    entry_slots[f"h_r{r}_c{c}"] = {"type": "harvey", "optional": True}
            # sample table text
            sample_rows = [[" "] + [f"[Criterion {j}]" for j in range(1, len(cols) + 1)]]
            sample_rows += [[f"[Option {i}]"] + [""] * len(cols) for i in range(1, len(rows) + 1)]
            dk.fill_table(tbl_shape, sample_rows)
            entry_slots["table"]["sample"] = sample_rows
            entry_slots["table"]["note"] = "First row: criteria; first column: options. Harvey levels go in h_r<row>_c<col> (0-4)."

        # --- risk stars: one per row, deleted for rows that are not top risks
        if spec.get("stars"):
            import copy as _copy
            proto = next((s for s in stars if s is not None), None)
            if proto is not None:
                for k in range(1, 8):
                    row = dk.slot_shapes(slide).get(f"r{k}_name")
                    if row is None:
                        continue
                    el = _copy.deepcopy(proto._element)
                    slide.shapes._spTree.append(el)
                    star = slide.shapes[-1]
                    star.top = row.top + (row.height - star.height) // 2
                    star.name = f"slot:r{k}_star"
                    entry_slots[f"r{k}_star"] = {"type": "marker", "optional": True}
                for s in stars:
                    if s is not None and s._element.getparent() is not None:
                        dk.remove_shape(s)

        if spec.get("add_source") and tajin_src is not None and "source" not in entry_slots:
            box = copy_source_box(slide, tajin_src.slides[1])
            if box is not None:
                dk.set_text(box, "Source: [Source]")
                box.name = "slot:source"
                entry_slots["source"] = {"type": "text", "sample": "Source: [Source]", "max_chars": 140, "font_pt": 9.0}

        # --- repeat groups: tag non-slot shapes that belong to each group
        repeats = []
        for rep in spec.get("repeat", []):
            prefix, count = rep["prefix"], rep["count"]
            axis = rep.get("axis", "x")
            members = {}
            for k in range(1, count + 1):
                anchors = []
                for slot in entry_slots:
                    if re.match(rf"^{prefix}{k}_", slot):
                        sh = dk.slot_shapes(slide).get(slot)
                        if sh is not None:
                            anchors.append(dk.top_level_ancestor(slide, sh))
                members[k] = anchors
            boxes = {k: dk.union(dk.bbox(a) for a in v) for k, v in members.items() if v}
            tol = int(0.15 * dk.IN)
            for sh in slide.shapes:
                if sh.name.startswith(("slot:", "grp:")) or sh.is_placeholder:
                    continue
                b = dk.bbox(sh)
                c = dk.center(b)
                for k, gb in boxes.items():
                    span = (gb[2] - gb[0]) if axis == "x" else (gb[3] - gb[1])
                    size = (b[2] - b[0]) if axis == "x" else (b[3] - b[1])
                    if dk.inside(c, gb, tol) and size <= span * 1.15:
                        sh.name = f"grp:{prefix}{k}:{sh.name}"
                        break
            repeats.append({"prefix": prefix, "count": count, "min": rep.get("min", 1),
                            "axis": axis, "mode": rep.get("mode", "move")})

        # tag the slide so inspect/review know the pattern
        tag = slide.shapes.add_textbox(0, 0, 0, 0)
        tag.name = f"pattern:{sid}"

        entry = {"purpose": spec["purpose"], "when_to_use": spec["when"],
                 "source": f"{fname} slide {num}", "layout": slide.slide_layout.name,
                 "slots": entry_slots}
        if repeats:
            entry["repeat"] = repeats
        if spec.get("internal"):
            entry["internal"] = True
        if sid == "framework":
            entry["tracker_mode"] = "Set `current: <k>` to brighten pillar k and keep the others shaded."
        catalog["patterns"][sid] = entry
        print(f"  {sid:18} {len(entry_slots):3} slots  from {fname} #{num}")

    colors, fonts = set(), set()
    for s in dst.slides:
        for el in s._element.iter():
            tag = str(el.tag)
            if tag.endswith("}srgbClr") and el.get("val"):
                colors.add(el.get("val").upper())
            if tag.endswith("}latin") and el.get("typeface") and not el.get("typeface").startswith("+"):
                fonts.add(el.get("typeface"))
    catalog["house_colors"] = sorted(colors)
    catalog["house_fonts"] = sorted(fonts)
    clear_authors(dst)
    for s in dst.slides:
        dk.drop_unused_rels(s)  # e.g. team photos replaced by a grey fill
    out_pptx = out_dir / "patterns.pptx"
    dst.save(str(out_pptx))
    with open(out_dir / "catalog.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump(catalog, f, sort_keys=False, allow_unicode=True, width=110)
    return out_pptx


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--winning", required=True, help="folder with the winning decks and the template")
    ap.add_argument("--out", default=str(dk.ASSETS))
    ap.add_argument("--only", nargs="*")
    args = ap.parse_args()
    winning = Path(args.winning)
    before = {p.name: p.stat().st_mtime for p in winning.glob("*.pptx")}
    out = curate(winning, Path(args.out), args.only)
    after = {p.name: p.stat().st_mtime for p in winning.glob("*.pptx")}
    assert before == after, "a source deck was modified!"
    print("wrote", out)


if __name__ == "__main__":
    main()
